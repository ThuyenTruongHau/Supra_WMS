from io import BytesIO

from openpyxl import Workbook

from app.modules.masan.masan_outbound_excel import DATA_START_ROW, HEADER_ROW
from app.modules.masan.masan_outbound_service import parse_masan_outbound_rows


def _build_minimal_bm04_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    headers = (
        "Số xe",
        "Tên khách hàng",
        "Trip",
        "NVT",
        "Mã Item",
        "Tên Item",
        "LOT",
        "Lot status",
        "Số lượng",
        "Số pallet",
        "Locator",
    )
    for col, title in enumerate(headers, start=1):
        ws.cell(HEADER_ROW, col, value=title)

    ws.cell(DATA_START_ROW, 1, value="29A-11111")
    ws.cell(DATA_START_ROW, 5, value="SKU-001")
    ws.cell(DATA_START_ROW, 7, value="220926")
    ws.cell(DATA_START_ROW, 9, value=5)
    ws.cell(DATA_START_ROW, 11, value="G101_RB1")

    ws.cell(DATA_START_ROW + 1, 5, value="SKU-002")
    ws.cell(DATA_START_ROW + 1, 9, value=2)
    ws.cell(DATA_START_ROW + 1, 11, value="G102_RB1")
    ws.merge_cells(
        start_row=DATA_START_ROW + 1,
        start_column=9,
        end_row=DATA_START_ROW + 2,
        end_column=9,
    )

    ws.cell(DATA_START_ROW + 2, 5, value="SKU-003")
    ws.cell(DATA_START_ROW + 2, 11, value="G103_RB1")

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


import unittest


class MasanOutboundParseTests(unittest.TestCase):
    def test_parse_resolves_merged_quantity_and_carry_fields(self):
        rows = parse_masan_outbound_rows(_build_minimal_bm04_workbook())
        by_sku = {row["sku"]: row for row in rows}

        self.assertEqual(len(rows), 3)
        self.assertEqual(by_sku["SKU-002"]["quantity"], 2)
        self.assertIsNone(by_sku["SKU-002"]["quantity_error"])
        self.assertEqual(by_sku["SKU-003"]["quantity"], 2)
        self.assertIsNone(by_sku["SKU-003"]["quantity_error"])
        self.assertEqual(by_sku["SKU-003"]["locator"], "G103_RB1")

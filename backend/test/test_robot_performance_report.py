"""Robot performance report: one count per robot task, SQLite only, no Redis/RCS."""

import os
import unittest
from datetime import datetime
from io import BytesIO

os.environ.update(
    DEBUG="false",
    DATABASE_URL="postgresql://unused@127.0.0.1:1/never_connect_in_tests",
    REDIS_URL="redis://127.0.0.1:1/0",
    CELERY_BROKER_URL="memory://",
    CELERY_BROKER_RESULT_URL="cache+memory://",
    APP_NAME="robot-report-test",
)

import app.main  # noqa: F401,E402  registers every model on Base.metadata
from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import create_engine, insert  # noqa: E402
from sqlalchemy.orm import Session  # noqa: E402
from sqlalchemy.pool import StaticPool  # noqa: E402

from app.core.database import Base  # noqa: E402
from app.modules.warehouse.dashboard import dashboard_service as svc  # noqa: E402

START = datetime(2026, 10, 23, 0, 0)
END = datetime(2026, 10, 23, 23, 59)


class RobotPerformanceReportTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite+pysqlite:///:memory:", poolclass=StaticPool,
                                    connect_args={"check_same_thread": False})
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self.addCleanup(self.engine.dispose)
        self.addCleanup(self.db.close)

        for wh in (1, 2):
            self.put("warehouse", id=wh, code=f"W{wh}")
        self.put("unit", id=1, name="Thung")
        self.put("users", id=1, username="test", email="test@example.invalid", hashed_password="unused")
        self.put("item", id=1, sku="SKU-1", name="Item", base_unit=1, max_quantity=9999,
                 min_quantity=0, warehouse_id=1, supplier="TEST")
        self.put("location", id=101, location_name="L1", location_code="L1", warehouse_id=1)
        self.put("item_stock", id=1, item_id=1, unit_id=1, location_id=101, quantity=5,
                 status="available", is_active=True)

        self.put("inbound_order", id=1, order_code="IN-1", created_by_id=1, warehouse_id=1)
        self.put("inbound_order_detail", id=1, inbound_order_id=1, status="completed",
                 detail_type="auto", distance=10.5)
        self.task(1, "TDS_Inbound_1", inbound_order_detail_id=1)

        for order_id, wh in ((1, 1), (2, 2)):
            self.put("outbound_order", id=order_id, order_code=f"OUT-{order_id}",
                     warehouse_id=wh, created_by_id=1)
            self.put("outbound_order_detail", id=order_id, outbound_order_id=order_id, item_id=1,
                     quantity=5, unit="Thung", detail_type="auto", details={})

        self.task(2, "TDS_Outbound_2")
        for ident in (1, 2, 3):
            self.allocation(ident, robot_task_id=2, detail_id=1, kind="outbound", distance=20.0)
        self.task(3, "TDS_Return_3")
        self.allocation(4, robot_task_id=3, detail_id=1, kind="return", distance=5.0)
        self.task(4, "TDS_Outbound_4")
        self.allocation(5, robot_task_id=4, detail_id=2, kind="outbound", distance=99.0)
        self.task(5, "TDS_Outbound_5")
        self.allocation(6, robot_task_id=5, detail_id=1, kind="outbound", distance=99.0)

        self.status("TDS_Inbound_1", "R1", "7", 8)
        self.status("TDS_Inbound_1", "R1", "8", 9)
        self.status("TDS_Outbound_2", "R1", "8", 10)
        self.status("TDS_Return_3", "R2", "7", 11)
        self.status("TDS_Outbound_4", "R1", "8", 12)  # other warehouse
        self.status("TDS_Outbound_5", "R1", "6", 13)  # not terminal
        self.db.commit()

    def put(self, table, **values):
        self.db.execute(insert(Base.metadata.tables[table]).values(**values))

    def task(self, ident, order_id, **values):
        self.put("robot_task", id=ident, order_id=order_id, quantity=1, process_code="P",
                 system_code="Thadosoft", task_order_detail="[]", **values)

    def allocation(self, ident, *, robot_task_id, detail_id, kind, distance):
        self.put("outbound_order_allocation", id=ident, outbound_order_detail_id=detail_id,
                 item_stock_id=1, quantity=1, status="completed", allocation_type=kind,
                 robot_task_id=robot_task_id, distance=distance)

    def status(self, order_id, device, status, hour):
        self.put("task_status", order_id=order_id, device_num=device, status=status,
                 created_at=datetime(2026, 10, 23, hour, 0))

    def test_counts_each_robot_task_once(self):
        rows = svc._query_robot_terminal_rows(self.db, warehouse_id=1, start=START, end=END)
        by_task = {row.robot_task_id: row for row in rows}
        self.assertEqual(sorted(by_task), [1, 2, 3])
        self.assertEqual(by_task[1].ics_status, "8")
        self.assertEqual(by_task[2].task_kind, "outbound")
        self.assertEqual(by_task[2].distance_m, 20.0)
        self.assertEqual(by_task[3].task_kind, "return")

        report = {r["device_num"]: r for r in svc._aggregate_robot_report(rows)}
        r1, r2, total = report["R1"], report["R2"], report["Tổng cộng"]
        self.assertEqual((r1["inbound_count"], r1["inbound_completed"], r1["inbound_error_count"]), (1, 1, 0))
        self.assertEqual((r1["outbound_count"], r1["outbound_completed"]), (1, 1))
        self.assertEqual(r1["distance_m"], 30.5)
        self.assertEqual((r2["return_count"], r2["return_error_count"], r2["distance_m"]), (1, 1, 0.0))
        self.assertEqual(total["total_count"], 3)
        self.assertEqual(total["distance_m"], 30.5)

    def test_workbook_row_matches_headers(self):
        rows = svc._aggregate_robot_report(
            svc._query_robot_terminal_rows(self.db, warehouse_id=1, start=START, end=END)
        )
        content = svc._build_robot_performance_workbook(
            period_label="x", warehouse_name="W1", rows=rows
        )
        ws = load_workbook(BytesIO(content)).active
        headers = [c.value for c in ws[svc._HEADER_ROW]]
        self.assertEqual(tuple(headers), svc._ROBOT_REPORT_HEADERS)
        r1 = [c.value for c in ws[svc._DATA_START_ROW]]
        self.assertEqual(r1, ["R1", 1, 1, 0, 2, 1, 1, 0, 0, 0, 0, 30.5])


if __name__ == "__main__":
    unittest.main()

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from io import BytesIO
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.modules.robot.robot_model import RobotTask, TaskStatus
from app.modules.warehouse.dashboard.dashboard_schema import ReportKpiResponse
from app.modules.warehouse.dashboard.dashboard_snapshot_service import VN_TZ, vn_today
from app.modules.warehouse.inbound_order import inbound_order_service
from app.modules.warehouse.inbound_order.inbound_order_model import (
    InboundOrder,
    InboundOrderDetail,
)
from app.modules.warehouse.item import item_service
from app.modules.warehouse.notificcation import notification_service
from app.modules.warehouse.outbound_order import outbound_order_service
from app.modules.warehouse.outbound_order.outbound_order_model import (
    OutboundOrder,
    OutboundOrderAllocation,
    OutboundOrderDetail,
)
from app.modules.warehouse.warehouse_zone.warehouse_model import Warehouse
from app.modules.warehouse.warehouse_zone.warehouse_service import _ensure_warehouse_exists

RobotReportPeriod = Literal["day", "week", "timeline"]
TaskKind = Literal["inbound", "outbound", "return"]

_TERMINAL_ICS_STATUSES = ("3", "7", "8")
_TASK_KINDS: tuple[TaskKind, ...] = ("inbound", "outbound", "return")

_ROBOT_REPORT_TITLE = (
    "BÁO CÁO HIỆU SUẤT ROBOT\n"
    "(Theo task_status ICS — mã 3/7/8, mỗi lệnh robot tính 1 lần)"
)
_ROBOT_REPORT_HEADERS: tuple[str, ...] = (
    "Robot",
    "Số lệnh robot nhập",
    "Số lệnh robot xuất",
    "Số lệnh robot hoàn trả",
    "Tổng số lệnh robot",
    "Số lệnh nhập hoàn thành",
    "Số lệnh xuất hoàn thành",
    "Số lệnh hoàn trả hoàn thành",
    "Số lệnh nhập lỗi",
    "Số lệnh xuất lỗi",
    "Số lệnh hoàn trả lỗi",
    "Tổng quãng đường (m)",
)
_ROBOT_REPORT_COLUMN_WIDTHS = (
    14.0, 16.0, 16.0, 18.0, 16.0, 20.0, 20.0, 22.0, 16.0, 16.0, 18.0, 18.0,
)
_HEADER_ROW = 3
_DATA_START_ROW = 4

_TITLE_FILL = PatternFill("solid", fgColor="FFFF00")
_HEADER_FONT = Font(bold=True, size=10)
_TITLE_FONT = Font(bold=True, size=11)
_DATA_FONT = Font(size=10)
_THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)


@dataclass(frozen=True)
class _RobotTerminalRow:
    device_num: str
    robot_task_id: int
    order_id: str
    task_kind: TaskKind
    ics_status: str
    distance_m: float | None = None


def get_report_kpis(db: Session, *, warehouse_id: int) -> ReportKpiResponse:
    _ensure_warehouse_exists(db, warehouse_id)
    return ReportKpiResponse(
        total_inbound_orders=inbound_order_service.count_inbound_orders(
            db, warehouse_id
        ),
        total_outbound_orders=outbound_order_service.count_outbound_orders(
            db, warehouse_id
        ),
        total_inventory_value=item_service.compute_total_inventory_value(
            db, warehouse_id
        ),
        unsolved_notifications=notification_service.count_unsolved_notifications(
            db, warehouse_id
        ),
    )


def _normalize_device_num(raw: str | None) -> str | None:
    if raw is None:
        return None
    cleaned = str(raw).strip()
    return cleaned or None


def _normalize_ics_status(raw: str | None) -> str | None:
    if raw is None:
        return None
    cleaned = str(raw).strip()
    return cleaned if cleaned in _TERMINAL_ICS_STATUSES else None


def _robot_report_period_bounds(
    period: RobotReportPeriod,
    *,
    start_date: date | None = None,
    end_date: date | None = None,
) -> tuple[datetime, datetime, str]:
    today = vn_today()
    if period == "day":
        start_date = today
        end_date = today
        label = f"Ngày {today.strftime('%d/%m/%Y')}"
    elif period == "week":
        start_date = today - timedelta(days=6)
        end_date = today
        label = (
            f"Từ {start_date.strftime('%d/%m/%Y')} "
            f"đến {end_date.strftime('%d/%m/%Y')}"
        )
    else:
        if start_date is None or end_date is None:
            raise ValueError("Timeline cần ngày bắt đầu và ngày kết thúc")
        if start_date > end_date:
            raise ValueError("Ngày bắt đầu không được sau ngày kết thúc")
        label = (
            f"Từ {start_date.strftime('%d/%m/%Y')} "
            f"đến {end_date.strftime('%d/%m/%Y')}"
        )
    start_dt = datetime.combine(start_date, time.min, tzinfo=VN_TZ)
    end_dt = datetime.combine(end_date, time.max.replace(microsecond=999999), tzinfo=VN_TZ)
    return start_dt, end_dt, label


def _max_distance(values: list[float | None]) -> float | None:
    present = [float(v) for v in values if v is not None]
    return max(present) if present else None


def _query_robot_terminal_rows(
    db: Session,
    *,
    warehouse_id: int,
    start: datetime,
    end: datetime,
) -> list[_RobotTerminalRow]:
    """One row per robot task (RobotTask.id): its latest terminal status 3/7/8 in the period.

    Allocations/details are only used to classify the task and read its distance, never to
    multiply the count: several allocations can share one robot task.
    """
    status_rows = (
        db.query(
            TaskStatus.device_num,
            TaskStatus.status,
            RobotTask.id.label("robot_task_id"),
            RobotTask.order_id,
            RobotTask.inbound_order_detail_id,
        )
        .join(RobotTask, TaskStatus.order_id == RobotTask.order_id)
        .filter(
            TaskStatus.created_at >= start,
            TaskStatus.created_at <= end,
            TaskStatus.status.in_(_TERMINAL_ICS_STATUSES),
        )
        .order_by(TaskStatus.created_at.desc(), TaskStatus.id.desc())
        .all()
    )

    latest: dict[int, tuple[str, str, str, int | None]] = {}
    for row in status_rows:
        if row.robot_task_id in latest:
            continue
        device_num = _normalize_device_num(row.device_num)
        ics_status = _normalize_ics_status(row.status)
        if not device_num or not ics_status:
            continue
        latest[row.robot_task_id] = (
            device_num,
            ics_status,
            (row.order_id or "").strip(),
            row.inbound_order_detail_id,
        )
    if not latest:
        return []

    inbound_detail_ids = {v[3] for v in latest.values() if v[3] is not None}
    inbound_info: dict[int, float | None] = {}
    if inbound_detail_ids:
        for detail_id, distance in (
            db.query(InboundOrderDetail.id, InboundOrderDetail.distance)
            .join(InboundOrder, InboundOrderDetail.inbound_order_id == InboundOrder.id)
            .filter(
                InboundOrderDetail.id.in_(inbound_detail_ids),
                InboundOrder.warehouse_id == warehouse_id,
            )
            .all()
        ):
            inbound_info[detail_id] = distance

    other_task_ids = [task_id for task_id, v in latest.items() if v[3] is None]
    allocation_types: dict[int, set[str]] = defaultdict(set)
    allocation_distances: dict[int, list[float | None]] = defaultdict(list)
    if other_task_ids:
        for task_id, allocation_type, distance in (
            db.query(
                OutboundOrderAllocation.robot_task_id,
                OutboundOrderAllocation.allocation_type,
                OutboundOrderAllocation.distance,
            )
            .join(
                OutboundOrderDetail,
                OutboundOrderAllocation.outbound_order_detail_id == OutboundOrderDetail.id,
            )
            .join(OutboundOrder, OutboundOrderDetail.outbound_order_id == OutboundOrder.id)
            .filter(
                OutboundOrderAllocation.robot_task_id.in_(other_task_ids),
                OutboundOrder.warehouse_id == warehouse_id,
            )
            .all()
        ):
            allocation_types[task_id].add(allocation_type)
            allocation_distances[task_id].append(distance)

    rows: list[_RobotTerminalRow] = []
    for task_id, (device_num, ics_status, order_id, detail_id) in latest.items():
        if detail_id is not None:
            if detail_id not in inbound_info:
                continue
            task_kind: TaskKind = "inbound"
            distance_m = inbound_info[detail_id]
        else:
            types = allocation_types.get(task_id)
            if not types:
                continue
            if "return" in types:
                task_kind = "return"
            elif "outbound" in types:
                task_kind = "outbound"
            else:
                continue
            distance_m = _max_distance(allocation_distances[task_id])
        rows.append(
            _RobotTerminalRow(
                device_num=device_num,
                robot_task_id=task_id,
                order_id=order_id,
                task_kind=task_kind,
                ics_status=ics_status,
                distance_m=distance_m,
            )
        )
    return rows


def _empty_robot_bucket() -> dict[str, float]:
    bucket: dict[str, float] = {"distance_m": 0.0}
    for kind in _TASK_KINDS:
        bucket[f"{kind}_count"] = 0
        bucket[f"{kind}_completed"] = 0
        bucket[f"{kind}_error_count"] = 0
    return bucket


def _aggregate_robot_report(
    rows: list[_RobotTerminalRow],
) -> list[dict[str, str | int | float]]:
    buckets: dict[str, dict[str, float]] = defaultdict(_empty_robot_bucket)

    for row in rows:
        bucket = buckets[row.device_num]
        bucket[f"{row.task_kind}_count"] += 1
        if row.ics_status == "8":
            bucket[f"{row.task_kind}_completed"] += 1
            if row.distance_m is not None:
                bucket["distance_m"] += row.distance_m
        elif row.ics_status in ("3", "7"):
            bucket[f"{row.task_kind}_error_count"] += 1

    def finalize(device_num: str, bucket: dict[str, float]) -> dict[str, str | int | float]:
        out: dict[str, str | int | float] = {"device_num": device_num}
        for key, value in bucket.items():
            out[key] = round(value, 3) if key == "distance_m" else int(value)
        out["total_count"] = sum(int(bucket[f"{kind}_count"]) for kind in _TASK_KINDS)
        return out

    result = [finalize(device, buckets[device]) for device in sorted(buckets)]
    if result:
        total = _empty_robot_bucket()
        for bucket in buckets.values():
            for key, value in bucket.items():
                total[key] += value
        result.append(finalize("Tổng cộng", total))
    return result


def _build_robot_performance_workbook(
    *,
    period_label: str,
    warehouse_name: str,
    rows: list[dict[str, str | int | float]],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Robot"

    col_count = len(_ROBOT_REPORT_HEADERS)
    for col_idx in range(1, col_count + 1):
        ws.cell(row=1, column=col_idx).fill = _TITLE_FILL

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = _ROBOT_REPORT_TITLE
    title_cell.font = _TITLE_FONT
    title_cell.alignment = _CENTER_WRAP
    ws.row_dimensions[1].height = 42.0

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    meta_cell = ws.cell(row=2, column=1)
    meta_cell.value = f"Kho: {warehouse_name} | Kỳ: {period_label}"
    meta_cell.font = Font(size=10, bold=True)
    meta_cell.alignment = _CENTER_WRAP
    ws.row_dimensions[2].height = 22.0

    for col_idx, header in enumerate(_ROBOT_REPORT_HEADERS, start=1):
        cell = ws.cell(row=_HEADER_ROW, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER_WRAP
        cell.border = _THIN_BORDER
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = _ROBOT_REPORT_COLUMN_WIDTHS[col_idx - 1]
    ws.row_dimensions[_HEADER_ROW].height = 28.0

    last_data_row = _DATA_START_ROW - 1
    for offset, row in enumerate(rows):
        row_idx = _DATA_START_ROW + offset
        last_data_row = row_idx
        values = [
            row["device_num"],
            row["inbound_count"],
            row["outbound_count"],
            row["return_count"],
            row["total_count"],
            row["inbound_completed"],
            row["outbound_completed"],
            row["return_completed"],
            row["inbound_error_count"],
            row["outbound_error_count"],
            row["return_error_count"],
            row["distance_m"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = _DATA_FONT
            cell.border = _THIN_BORDER
            if col_idx == 1:
                cell.alignment = _LEFT_WRAP
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
        if row["device_num"] == "Tổng cộng":
            for col_idx in range(1, col_count + 1):
                ws.cell(row=row_idx, column=col_idx).font = Font(size=10, bold=True)

    for row_idx in range(_DATA_START_ROW, last_data_row + 1):
        ws.row_dimensions[row_idx].height = 20.0

    ws.freeze_panes = ws.cell(row=_DATA_START_ROW, column=1)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_robot_performance_report(
    db: Session,
    *,
    warehouse_id: int,
    period: RobotReportPeriod,
    start_date: date | None = None,
    end_date: date | None = None,
) -> tuple[bytes, str]:
    _ensure_warehouse_exists(db, warehouse_id)
    warehouse = db.query(Warehouse).filter(Warehouse.id == warehouse_id).first()
    warehouse_label = (
        (warehouse.code if warehouse and warehouse.code else None)
        or (warehouse.name if warehouse and warehouse.name else None)
        or str(warehouse_id)
    )
    warehouse_name = warehouse.name if warehouse and warehouse.name else warehouse_label

    start_dt, end_dt, period_label = _robot_report_period_bounds(
        period,
        start_date=start_date,
        end_date=end_date,
    )
    terminal_rows = _query_robot_terminal_rows(
        db,
        warehouse_id=warehouse_id,
        start=start_dt,
        end=end_dt,
    )
    if not terminal_rows:
        raise ValueError(
            "Không có lệnh robot nào có trạng thái 3/7/8 trong kỳ đã chọn"
        )

    aggregated = _aggregate_robot_report(terminal_rows)
    content = _build_robot_performance_workbook(
        period_label=period_label,
        warehouse_name=warehouse_name,
        rows=aggregated,
    )
    if period == "day":
        period_suffix = f"ngay_{vn_today().isoformat()}"
    elif period == "week":
        period_suffix = f"tuan_{vn_today().isoformat()}"
    else:
        if start_date is None or end_date is None:
            raise ValueError("Timeline cần ngày bắt đầu và ngày kết thúc")
        period_suffix = f"timeline_{start_date.isoformat()}_{end_date.isoformat()}"
    filename = f"BaoCaoRobot_{warehouse_label}_{period_suffix}.xlsx"
    return content, filename

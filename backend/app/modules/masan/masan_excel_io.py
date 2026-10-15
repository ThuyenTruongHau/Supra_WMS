"""Centralized openpyxl load for Masan templates (suppress benign extension warnings)."""

from __future__ import annotations

import warnings
from io import BytesIO
from pathlib import Path
from typing import BinaryIO, Union

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

FileSource = Union[bytes, BinaryIO, str, Path]

_OPENPYXL_READER = "openpyxl.worksheet._reader"


def _suppress_openpyxl_extension_warnings() -> None:
    warnings.filterwarnings(
        "ignore",
        message=r".*Slicer List extension is not supported.*",
        category=UserWarning,
        module=_OPENPYXL_READER,
    )
    warnings.filterwarnings(
        "ignore",
        message=r".*Data Validation extension is not supported.*",
        category=UserWarning,
        module=_OPENPYXL_READER,
    )


def load_masan_workbook(source: FileSource, *, data_only: bool = False) -> Workbook:
    """
    Load BM.04 / Masan xlsx. Templates from Excel often include Slicer metadata;
    openpyxl drops it safely — no need to spam uvicorn logs.
    """
    with warnings.catch_warnings():
        _suppress_openpyxl_extension_warnings()
        if isinstance(source, bytes):
            return load_workbook(BytesIO(source), data_only=data_only)
        if isinstance(source, (str, Path)):
            return load_workbook(source, data_only=data_only)
        return load_workbook(source, data_only=data_only)

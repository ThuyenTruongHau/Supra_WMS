"""Shared Excel layout for Masan inbound import/export (BM.06)."""

from __future__ import annotations

from typing import TYPE_CHECKING, Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

EXPORT_TITLE = (
    "BIỂU MẪU BÁO CÁO NHẬP THEO NGÀY\n(New System trả về để truy vết)"
)

EXPORT_HEADERS: tuple[str, ...] = (
    "Ngày/ Giờ chi tiết",
    "Số xe",
    "Kho xuất",
    "Kho nhập",
    "Delivery",
    "NVT",
    "Mã Item",
    "Tên Item",
    "LOT",
    "Lot status",
    "Số lượng",
    "Số pallet",
    "Vị trí để",
    "Locator",
)

INBOUND_REPORT_EXTRA_HEADERS: tuple[str, ...] = (
    "Người nhập",
    "Quãng đường (m)",
    "Robot",
    "Thời gian bắt đầu",
    "Thời gian kết thúc",
    "Pin bắt đầu",
    "Pin kết thúc",
)

INBOUND_REPORT_HEADERS: tuple[str, ...] = EXPORT_HEADERS + INBOUND_REPORT_EXTRA_HEADERS

HEADER_ROW = 2
DATA_START_ROW = 3

INBOUND_REPORT_COLUMN_WIDTHS: tuple[float, ...] = (
    14.0,  # Ngày/ Giờ chi tiết
    12.0,  # Số xe
    30.0,  # Kho xuất
    30.0,  # Kho nhập
    14.0,  # Delivery
    14.0,  # NVT
    12.0,  # Mã Item
    45.0,  # Tên Item
    10.0,  # LOT
    12.0,  # Lot status
    10.0,  # Số lượng
    10.0,  # Số pallet
    14.0,  # Vị trí để
    14.0,  # Locator
    18.0,  # Người nhập
    12.0,  # Quãng đường (m)
    10.0,  # Robot
    18.0,  # Thời gian bắt đầu
    18.0,  # Thời gian kết thúc
    12.0,  # Pin bắt đầu
    12.0,  # Pin kết thúc
)

TITLE_ROW_HEIGHT = 42.0
HEADER_ROW_HEIGHT = 28.0
DATA_ROW_HEIGHT = 20.0

# Columns that must stay text (avoid scientific notation / leading-zero loss).
TEXT_FORMAT_COLUMNS = frozenset({5, 9})  # Delivery, LOT
INBOUND_REPORT_QUANTITY_COLUMN = 11
INBOUND_REPORT_DISTANCE_COLUMN = 16
WRAP_TEXT_COLUMNS = frozenset({1, 8, 15, 18, 19})  # Ngày/Giờ, Tên Item, người nhập, thời gian robot

_TITLE_FILL = PatternFill("solid", fgColor="FFFF00")
_HEADER_FONT = Font(bold=True, size=10)
_TITLE_FONT = Font(bold=True, size=11)
_DATA_FONT = Font(size=10)
_THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
_LEFT_TOP_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def normalize_export_cell_value(column_index: int, value: Any) -> Any:
    """Coerce values for stable Excel display (text IDs, preserved blanks)."""
    if value is None or value == "":
        return None
    if column_index in TEXT_FORMAT_COLUMNS:
        return str(value)
    return value


def apply_export_cell_style(ws: Worksheet, row: int, column: int) -> None:
    cell = ws.cell(row=row, column=column)
    cell.font = _DATA_FONT
    cell.border = _THIN_BORDER
    if column in TEXT_FORMAT_COLUMNS:
        cell.number_format = "@"
    if column in WRAP_TEXT_COLUMNS:
        cell.alignment = _LEFT_TOP_WRAP
    elif column in (INBOUND_REPORT_QUANTITY_COLUMN, INBOUND_REPORT_DISTANCE_COLUMN):
        cell.alignment = Alignment(horizontal="right", vertical="center")
    else:
        cell.alignment = _LEFT_WRAP


def style_masan_export_worksheet(ws: Worksheet, *, last_data_row: int) -> None:
    """Apply BM.06-like layout: yellow title row, spaced columns/rows, grid borders."""
    col_count = len(INBOUND_REPORT_HEADERS)

    for col_idx, width in enumerate(INBOUND_REPORT_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for col_idx in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _TITLE_FILL

    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=col_count)
    title_cell = ws.cell(row=1, column=2)
    title_cell.value = EXPORT_TITLE
    title_cell.font = _TITLE_FONT
    title_cell.alignment = _CENTER_WRAP
    ws.row_dimensions[1].height = TITLE_ROW_HEIGHT

    for col_idx in range(1, col_count + 1):
        header_cell = ws.cell(row=HEADER_ROW, column=col_idx)
        header_cell.font = _HEADER_FONT
        header_cell.alignment = _CENTER_WRAP
        header_cell.border = _THIN_BORDER
    ws.row_dimensions[HEADER_ROW].height = HEADER_ROW_HEIGHT

    for row_idx in range(DATA_START_ROW, last_data_row + 1):
        ws.row_dimensions[row_idx].height = DATA_ROW_HEIGHT
        for col_idx in range(1, col_count + 1):
            apply_export_cell_style(ws, row_idx, col_idx)

    ws.freeze_panes = ws.cell(row=DATA_START_ROW, column=1)


_DAILY_CAPACITY_SUMMARY_COLS = 2
_DAILY_CAPACITY_TABLE_COLS = 7
_DAILY_CAPACITY_LOT_COLUMN = 6
_DAILY_CAPACITY_QTY_COLUMN = 5
_DAILY_SUMMARY_LABEL_FILL = PatternFill("solid", fgColor="E2EFDA")
_DAILY_TABLE_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")


def style_daily_capacity_worksheet(
    ws: Worksheet,
    *,
    summary_header_row: int,
    summary_last_row: int,
    table_header_row: int,
    last_data_row: int,
) -> None:
    """Summary (2 cols) + locator table (Locator, Trạng thái, Mã/Tên hàng, SL, LOT, Ngày nhập)."""
    table_width = _DAILY_CAPACITY_TABLE_COLS
    col_widths = {
        1: 14.0,
        2: 14.0,
        3: 14.0,
        4: 22.0,
        5: 10.0,
        6: 12.0,
        7: 14.0,
    }
    for col_idx in range(1, table_width + 1):
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = col_widths.get(col_idx, 12.0)

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=table_width,
    )
    title_cell = ws.cell(row=1, column=1)
    title_cell.fill = _TITLE_FILL
    title_cell.font = _TITLE_FONT
    title_cell.alignment = _CENTER_WRAP
    ws.row_dimensions[1].height = 36.0

    for row_idx in range(summary_header_row, summary_last_row + 1):
        ws.row_dimensions[row_idx].height = 22.0
        for col_idx in range(1, _DAILY_CAPACITY_SUMMARY_COLS + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = _THIN_BORDER
            cell.font = _HEADER_FONT if row_idx == summary_header_row else _DATA_FONT
            cell.alignment = _CENTER_WRAP if col_idx == 2 else _LEFT_WRAP
            if row_idx == summary_header_row:
                cell.fill = _DAILY_SUMMARY_LABEL_FILL
            elif col_idx == 1:
                cell.fill = _DAILY_SUMMARY_LABEL_FILL

    ws.row_dimensions[table_header_row].height = 24.0
    for col_idx in range(1, table_width + 1):
        cell = ws.cell(row=table_header_row, column=col_idx)
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER_WRAP
        cell.border = _THIN_BORDER
        cell.fill = _DAILY_TABLE_HEADER_FILL

    for row_idx in range(table_header_row + 1, last_data_row + 1):
        ws.row_dimensions[row_idx].height = 18.0
        for col_idx in range(1, table_width + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = _DATA_FONT
            cell.border = _THIN_BORDER
            if col_idx == 2:
                cell.alignment = _CENTER_WRAP
            elif col_idx == _DAILY_CAPACITY_QTY_COLUMN:
                cell.alignment = _CENTER_WRAP
            else:
                cell.alignment = _LEFT_TOP_WRAP
            if col_idx == _DAILY_CAPACITY_LOT_COLUMN:
                cell.number_format = "@"

    ws.freeze_panes = ws.cell(row=table_header_row + 1, column=1)


COLUMN_ALIASES: dict[str, list[str]] = {
    "inbound_datetime": ["ngày/ giờ chi tiết", "ngày/giờ chi tiết"],
    "vehicle_no": ["số xe"],
    "from_warehouse": ["kho xuất"],
    "to_warehouse": ["kho nhập"],
    "delivery": ["delivery"],
    "nvt": ["nvt"],
    "sku": ["mã item"],
    "item_name": ["tên item"],
    "lot": ["lot"],
    "lot_status": ["lot status"],
    "quantity": ["số lượng"],
    "pallet_count": ["số pallet"],
    "storage_location": ["vị trí để"],
    "locator": ["locator"],
}

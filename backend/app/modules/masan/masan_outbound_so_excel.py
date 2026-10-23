"""Excel layout for Masan outbound SO export (multi-sheet by customer)."""

from __future__ import annotations

import re
from typing import TYPE_CHECKING, Any

from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

SO_TITLE = "BIỂU MẪU LẤY HÀNG THEO SO"
SO_SUBTITLE = "(Cung cấp thông tin xuất hàng)"

SO_DATA_HEADERS: tuple[str, ...] = (
    "STT",
    "Mã Item",
    "Tên Item",
    "LOT",
    "Lot status",
    "Số lượng",
    "Số pallet",
    "Locator",
)

TITLE_ROW = 1
SUBTITLE_ROW = 2
CUSTOMER_ROW = 3
ADDRESS_ROW = 4
CARRIER_ROW = 5
HEADER_ROW = 6
DATA_START_ROW = 7

COL_COUNT = len(SO_DATA_HEADERS)

SO_COLUMN_WIDTHS: tuple[float, ...] = (
    7.8,
    14.8,
    52.8,
    11.8,
    13.8,
    11.8,
    11.8,
    14.8,
)

TITLE_ROW_HEIGHT = 30.0
SUBTITLE_ROW_HEIGHT = 22.0
META_ROW_HEIGHT = 28.0
HEADER_ROW_HEIGHT = 28.0
DATA_ROW_HEIGHT = 42.0

TEXT_FORMAT_COLUMNS = frozenset({4, 8})  # LOT, Locator
WRAP_TEXT_COLUMNS = frozenset({3, 8})  # Tên Item, Locator

_HEADER_FONT = Font(bold=True, size=12)
_TITLE_FONT = Font(bold=True, size=14)
_SUBTITLE_FONT = Font(bold=True, size=12)
_META_FONT = Font(bold=True, size=11)
_DATA_FONT = Font(bold=True, size=11)
_THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)

_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")


def sanitize_sheet_name(name: str, *, used: set[str] | None = None) -> str:
    cleaned = _INVALID_SHEET_CHARS.sub("", (name or "").strip()) or "Unknown"
    cleaned = cleaned[:31]
    if used is None:
        return cleaned
    base = cleaned
    candidate = base
    suffix = 2
    while candidate in used:
        tail = f"_{suffix}"
        candidate = f"{base[: 31 - len(tail)]}{tail}"
        suffix += 1
    used.add(candidate)
    return candidate


def normalize_so_cell_value(column_index: int, value: Any) -> Any:
    if value is None or value == "":
        return None
    if column_index in TEXT_FORMAT_COLUMNS:
        return str(value)
    return value


def _apply_data_cell_style(ws: Worksheet, row: int, column: int) -> None:
    cell = ws.cell(row=row, column=column)
    cell.font = _DATA_FONT
    cell.border = _THIN_BORDER
    if column in TEXT_FORMAT_COLUMNS:
        cell.number_format = "@"
    cell.alignment = _CENTER_WRAP


def build_so_customer_sheet(
    ws: Worksheet,
    *,
    customer_name: str,
    trips: list[str],
    vehicles: list[str],
    nvts: list[str],
    data_rows: list[list[Any]],
) -> None:
    for col_idx, width in enumerate(SO_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=COL_COUNT)
    ws.cell(row=TITLE_ROW, column=1, value=SO_TITLE).font = _TITLE_FONT
    ws.cell(row=TITLE_ROW, column=1).alignment = _CENTER_WRAP
    ws.row_dimensions[TITLE_ROW].height = TITLE_ROW_HEIGHT

    ws.merge_cells(
        start_row=SUBTITLE_ROW, start_column=1, end_row=SUBTITLE_ROW, end_column=COL_COUNT
    )
    ws.cell(row=SUBTITLE_ROW, column=1, value=SO_SUBTITLE).font = _SUBTITLE_FONT
    ws.cell(row=SUBTITLE_ROW, column=1).alignment = _CENTER_WRAP
    ws.row_dimensions[SUBTITLE_ROW].height = SUBTITLE_ROW_HEIGHT

    ws.merge_cells(start_row=CUSTOMER_ROW, start_column=1, end_row=CUSTOMER_ROW, end_column=4)
    customer_cell = ws.cell(row=CUSTOMER_ROW, column=1, value=f"Khách hàng: {customer_name}")
    customer_cell.font = _META_FONT
    customer_cell.alignment = _CENTER_WRAP
    ws.merge_cells(start_row=CUSTOMER_ROW, start_column=5, end_row=CUSTOMER_ROW, end_column=COL_COUNT)
    trip_text = ", ".join(trips) if trips else "…"
    trip_cell = ws.cell(row=CUSTOMER_ROW, column=5, value=f"Số đơn hàng: {trip_text}")
    trip_cell.font = _META_FONT
    trip_cell.alignment = _CENTER_WRAP
    ws.row_dimensions[CUSTOMER_ROW].height = META_ROW_HEIGHT

    ws.merge_cells(start_row=ADDRESS_ROW, start_column=1, end_row=ADDRESS_ROW, end_column=4)
    address_cell = ws.cell(row=ADDRESS_ROW, column=1, value="Địa chỉ giao hàng: —")
    address_cell.font = _META_FONT
    address_cell.alignment = _CENTER_WRAP
    ws.row_dimensions[ADDRESS_ROW].height = META_ROW_HEIGHT

    ws.merge_cells(start_row=CARRIER_ROW, start_column=1, end_row=CARRIER_ROW, end_column=4)
    nvt_text = ", ".join(nvts) if nvts else "—"
    carrier_cell = ws.cell(row=CARRIER_ROW, column=1, value=f"Đơn vị vận chuyển: {nvt_text}")
    carrier_cell.font = _META_FONT
    carrier_cell.alignment = _CENTER_WRAP
    ws.merge_cells(start_row=CARRIER_ROW, start_column=5, end_row=CARRIER_ROW, end_column=COL_COUNT)
    vehicle_text = ", ".join(vehicles) if vehicles else "—"
    vehicle_cell = ws.cell(row=CARRIER_ROW, column=5, value=f"Số xe: {vehicle_text}")
    vehicle_cell.font = _META_FONT
    vehicle_cell.alignment = _CENTER_WRAP
    ws.row_dimensions[CARRIER_ROW].height = META_ROW_HEIGHT

    for col_idx, header in enumerate(SO_DATA_HEADERS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER_WRAP
        cell.border = _THIN_BORDER
    ws.row_dimensions[HEADER_ROW].height = HEADER_ROW_HEIGHT

    for offset, row_values in enumerate(data_rows):
        row_idx = DATA_START_ROW + offset
        ws.row_dimensions[row_idx].height = DATA_ROW_HEIGHT
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(
                row=row_idx,
                column=col_idx,
                value=normalize_so_cell_value(col_idx, value),
            )
            _apply_data_cell_style(ws, row_idx, col_idx)

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_area = f"A1:{get_column_letter(COL_COUNT)}{HEADER_ROW + len(data_rows)}"
    ws.print_title_rows = f"1:{HEADER_ROW}"
    # Reserve a separate signature band so long tables cannot overlap the footer.
    ws.page_margins = PageMargins(
        left=0.3, right=0.3, top=0.4, bottom=1.45, header=0.15, footer=0.3
    )
    ws.HeaderFooter.scaleWithDoc = False
    ws.HeaderFooter.alignWithMargins = True
    ws.oddFooter.right.font = "Calibri"
    ws.oddFooter.right.size = 12
    ws.oddFooter.right.text = (
        "&BNgười chia hàng&B\n"
        "(Ký, ghi rõ họ tên)\n\n\n\n"
        "................................................"
    )

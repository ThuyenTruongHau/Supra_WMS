from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook
import re

from app.modules.masan.masan_excel_io import load_masan_workbook
from sqlalchemy import Integer, cast, func, or_
from sqlalchemy.orm import Session, joinedload, selectinload
from app.modules.warehouse.inbound_order.inbound_celery_task import accept_inbound_task_task
from sqlalchemy import desc
from app.core.config import settings
from app.core.cache import cache_set, cache_get

from app.modules.masan.masan_inbound_excel import (
    COLUMN_ALIASES,
    DATA_START_ROW,
    INBOUND_REPORT_HEADERS,
    HEADER_ROW,
    normalize_export_cell_value,
    style_daily_capacity_worksheet,
    style_masan_export_worksheet,
)
from app.modules.masan.masan_schema import (
    MasanInboundParseResponse,
    MasanInboundPreviewRow,
)
from app.modules.warehouse.inbound_order.inbound_order_model import (
    InboundOrder,
    InboundOrderAllocation,
    InboundOrderDetail,
)
from app.modules.warehouse.inbound_order.inbound_order_schema import (
    InboundAssignedDetailItem,
    InboundAssignedDetailsResponse,
    InboundOrderDetailResponse,
    InboundSuggestAdditionalDetail,
    InboundSuggestAllocation,
    InboundSuggestAllocationDetail,
)
from app.modules.warehouse.lot_number_utils import (
    apply_lot_display_fields,
    coerce_masan_lot_cell,
    format_lot_for_masan_excel,
    format_lot_number_display,
    format_lot_value_for_display,
    normalize_stored_lot_part,
)
from app.modules.warehouse.inbound_order.inbound_order_service import (
    _build_detail_response,
)
from app.modules.warehouse.item.item_model import Item
from app.modules.warehouse.item_stock.item_stock_model import (
    ItemStock,
    countable_stock_level_criterion,
    positive_stock_quantity_criterion,
)
from app.modules.warehouse.location_map.location_display import location_display_label
from app.modules.warehouse.location_map.location_model import Location
from app.modules.warehouse.location_map.location_service import (
    location_ids_with_assigned_qr,
    overlay_status_with_assign_cache,
)
from app.modules.warehouse.warehouse_zone.warehouse_model import Warehouse, Zone
from app.modules.warehouse.warehouse_zone.warehouse_service import _ensure_warehouse_exists
from app.core.logger import get_logger

logger = get_logger("main")


DELIVERY_PLACEHOLDERS = frozenset({"…", "...", ".", "-", "—"})


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower()


def _build_column_map(header_row: list[Any]) -> dict[str, int]:
    col_map: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        name = _normalize_header(cell)
        for field, aliases in COLUMN_ALIASES.items():
            if name in aliases:
                col_map[field] = idx
    return col_map


def _direct(row: list[Any], col_map: dict[str, int], field: str) -> Optional[str]:
    idx = col_map.get(field)
    if idx is None or idx >= len(row) or row[idx] in (None, ""):
        return None
    return str(row[idx]).strip()


def _normalize_delivery(value: Optional[str], carry: Optional[str]) -> Optional[str]:
    if value is None:
        return carry
    cleaned = value.strip()
    if not cleaned or cleaned in DELIVERY_PLACEHOLDERS:
        return carry
    return cleaned


def _cell_value(
    row: list[Any],
    col_map: dict[str, int],
    field: str,
    carry: dict[str, str],
) -> Optional[str]:
    idx = col_map.get(field)
    if idx is None:
        return carry.get(field)

    raw = row[idx] if idx < len(row) else None
    if raw is None or str(raw).strip() == "":
        return carry.get(field)

    value = str(raw).strip()
    if field == "delivery":
        normalized = _normalize_delivery(value, carry.get(field))
        if normalized:
            carry[field] = normalized
        return normalized

    carry[field] = value
    return value


def _parse_quantity(raw: Any) -> tuple[int, Optional[str]]:
    if raw in (None, ""):
        return 0, "Số lượng trống"
    try:
        qty = int(float(str(raw).replace(",", "")))
    except (TypeError, ValueError):
        return 0, f"Số lượng không hợp lệ: {raw!r}"
    if qty <= 0:
        return qty, "Số lượng phải lớn hơn 0"
    return qty, None


def parse_masan_inbound_rows(content: bytes) -> list[dict[str, Any]]:
    wb = load_masan_workbook(content, data_only=True)
    ws = wb.active

    header_row = [ws.cell(HEADER_ROW, c).value for c in range(1, ws.max_column + 1)]
    col_map = _build_column_map(header_row)

    if "sku" not in col_map:
        raise ValueError('Không tìm thấy cột "Mã Item" ở hàng 2')

    rows: list[dict[str, Any]] = []
    carry: dict[str, str] = {}

    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        row = [ws.cell(row_idx, c).value for c in range(1, ws.max_column + 1)]

        sku_raw = row[col_map["sku"]] if col_map["sku"] < len(row) else None
        if sku_raw is None or str(sku_raw).strip() == "":
            continue

        qty_idx = col_map.get("quantity")
        quantity_raw = row[qty_idx] if qty_idx is not None and qty_idx < len(row) else None
        quantity, qty_error = _parse_quantity(quantity_raw)

        rows.append(
            {
                "row_no": row_idx,
                "inbound_datetime": _cell_value(row, col_map, "inbound_datetime", carry),
                "vehicle_no": _cell_value(row, col_map, "vehicle_no", carry),
                "from_warehouse": _cell_value(row, col_map, "from_warehouse", carry),
                "to_warehouse": _cell_value(row, col_map, "to_warehouse", carry),
                "delivery": _cell_value(row, col_map, "delivery", carry),
                "nvt": _cell_value(row, col_map, "nvt", carry),
                "sku": str(sku_raw).strip(),
                "item_name": _direct(row, col_map, "item_name"),
                "lot": coerce_masan_lot_cell(_direct(row, col_map, "lot")),
                "lot_status": _direct(row, col_map, "lot_status"),
                "quantity": quantity,
                "quantity_error": qty_error,
                "pallet_count": _direct(row, col_map, "pallet_count"),
                "storage_location": _direct(row, col_map, "storage_location"),
                "locator": _direct(row, col_map, "locator"),
            }
        )

    if not rows:
        raise ValueError("Không có dòng hợp lệ trong file Excel")

    return rows


def _build_line_item_details(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "inbound_datetime": row.get("inbound_datetime"),
        "vehicle_no": row.get("vehicle_no"),
        "from_warehouse": row.get("from_warehouse"),
        "to_warehouse": row.get("to_warehouse"),
        "delivery": row.get("delivery"),
        "nvt": row.get("nvt"),
        "sku": row.get("sku"),
        "item_name": row.get("item_name"),
        "lot": coerce_masan_lot_cell(row.get("lot")),
        "lot_status": row.get("lot_status"),
        "pallet_count": row.get("pallet_count"),
        "storage_location": row.get("storage_location"),
        "from_location_id": row.get("from_location_id"),
        "from_location_name": row.get("from_location_name"),
        "locator": row.get("locator"),
    }


def _lookup_item(db: Session, warehouse_id: int, sku: str) -> Optional[Item]:
    return (
        db.query(Item)
        .filter(
            Item.warehouse_id == warehouse_id,
            Item.sku == sku,
            Item.is_active.is_(True),
        )
        .first()
    )


def _lookup_inbound_from_location(
    db: Session,
    warehouse_id: int,
    location_name: Optional[str],
) -> tuple[Optional[Location], Optional[str]]:
    if not location_name or not str(location_name).strip():
        return None, "Thiếu vị trí để"

    name = str(location_name).strip()
    zone_filter = or_(
        Zone.name.in_(settings.zone_inbound),
        Zone.code.in_(settings.zone_inbound),
    )

    matches = (
        db.query(Location)
        .join(Zone, Location.zone_id == Zone.id)
        .filter(
            Location.warehouse_id == warehouse_id,
            Location.is_active.is_(True),
            Location.location_name == name,
            zone_filter,
        )
        .all()
    )

    if not matches:
        exists_elsewhere = (
            db.query(Location)
            .filter(
                Location.warehouse_id == warehouse_id,
                Location.is_active.is_(True),
                Location.location_name == name,
            )
            .first()
        )
        if exists_elsewhere:
            return None, f"Vị trí '{name}' không thuộc zone nhập"
        return None, f"Không tìm thấy vị trí '{name}' trong kho"

    if len(matches) > 1:
        return None, f"Vị trí '{name}' không xác định (trùng tên trong zone nhập)"

    return matches[0], None


def parse_masan_inbound_preview(
    db: Session,
    *,
    warehouse_id: int,
    content: bytes,
    inbound_type: str,
) -> MasanInboundParseResponse:
    _ensure_warehouse_exists(db, warehouse_id)

    parsed_rows = parse_masan_inbound_rows(content)
    sku_cache: dict[str, Optional[Item]] = {}

    preview_rows: list[MasanInboundPreviewRow] = []
    line_items: list[InboundSuggestAdditionalDetail] = []
    suggest_index_by_row_no: dict[int, int] = {}
    warnings: list[str] = []
    invalid_count = 0

    for row in parsed_rows:
        sku = row["sku"]
        if sku not in sku_cache:
            sku_cache[sku] = _lookup_item(db, warehouse_id, sku)

        item = sku_cache[sku]
        from_location, from_location_error = _lookup_inbound_from_location(
            db,
            warehouse_id,
            row.get("storage_location"),
        )

        errors: list[str] = []
        if row.get("quantity_error"):
            errors.append(row["quantity_error"])
        if item is None:
            errors.append(f"Không tìm thấy sản phẩm SKU '{sku}' trong kho")
        if from_location_error:
            errors.append(from_location_error)

        preview = MasanInboundPreviewRow(
            row_no=row["row_no"],
            inbound_datetime=row.get("inbound_datetime"),
            vehicle_no=row.get("vehicle_no"),
            from_warehouse=row.get("from_warehouse"),
            to_warehouse=row.get("to_warehouse"),
            delivery=row.get("delivery"),
            nvt=row.get("nvt"),
            sku=sku,
            item_name=row.get("item_name"),
            lot=row.get("lot"),
            lot_status=row.get("lot_status"),
            quantity=row.get("quantity") or 0,
            pallet_count=row.get("pallet_count"),
            storage_location=row.get("storage_location"),
            from_location_id=from_location.id if from_location else None,
            from_location_name=from_location.location_name if from_location else None,
            locator=row.get("locator"),
            item_id=item.id if item else None,
            unit_id=item.base_unit_id if item else None,
            error="; ".join(errors) if errors else None,
        )
        preview_rows.append(preview)

        if errors or item is None or from_location is None:
            invalid_count += 1
            continue

        row_with_from_location = {
            **row,
            "from_location_id": from_location.id,
            "from_location_name": from_location.location_name,
        }
        suggest_index_by_row_no[row["row_no"]] = len(line_items)
        line_items.append(
            InboundSuggestAdditionalDetail(
                items=[
                    InboundSuggestAllocationDetail(
                        item_id=item.id,
                        quantity=row["quantity"],
                        unit_id=item.base_unit_id,
                        lot_number=row.get("lot"),
                    )
                ],
                details=_build_line_item_details(row_with_from_location),
            )
        )

    if not line_items:
        raise ValueError(
            "Không có dòng hợp lệ để tạo payload gợi ý vị trí. "
            "Kiểm tra SKU và số lượng trong file."
        )

    for idx, preview in enumerate(preview_rows):
        group_idx = suggest_index_by_row_no.get(preview.row_no)
        if group_idx is None:
            continue
        preview_rows[idx] = preview.model_copy(update={"suggest_group_index": group_idx})

    if invalid_count:
        warnings.append(f"{invalid_count} dòng bị bỏ qua khỏi payload gợi ý vị trí")

    suggest_allocation = InboundSuggestAllocation(
        warehouse_id=warehouse_id,
        detail_type=inbound_type,
        line_items=line_items,
    )
    for detail in suggest_allocation.line_items:
        for item in detail.items:
            disp_from, disp_to, disp_lot = apply_lot_display_fields(
                lot_number_from=item.lot_number_from,
                lot_number_to=item.lot_number_to,
                lot_number=item.lot_number,
            )
            item.lot_number_from = disp_from
            item.lot_number_to = disp_to
            item.lot_number = disp_lot
    for row in preview_rows:
        if row.lot:
            row.lot = format_lot_value_for_display(row.lot) or row.lot

    return MasanInboundParseResponse(
        preview_rows=preview_rows,
        suggest_allocation=suggest_allocation,
        total_rows=len(preview_rows),
        valid_rows=len(preview_rows) - invalid_count,
        invalid_rows=invalid_count,
        group_count=len(line_items),
        warnings=warnings,
    )


def _meta_value(meta: dict[str, Any], key: str, default: Any = "") -> Any:
    value = meta.get(key)
    if value is None or value == "":
        return default
    return value


def _format_inbound_datetime(order: InboundOrder, meta: dict[str, Any]) -> str:
    raw = _meta_value(meta, "inbound_datetime", default=None)
    if raw:
        return str(raw)
    if order.created_at:
        return order.created_at.strftime("%d/%m/%Y\n%H:%M")
    return ""


def _format_inbound_report_datetime(dt) -> str:
    if not dt:
        return ""
    return dt.strftime("%d/%m/%Y\n%H:%M")


def _build_export_row(
    order: InboundOrder,
    detail: InboundOrderDetail,
    allocation: InboundOrderAllocation,
) -> list[Any]:
    meta = detail.details or {}
    stock: ItemStock | None = allocation.item_stock
    item = stock.item if stock else None

    sku = item.sku if item else _meta_value(meta, "sku", default="")
    item_name = item.name if item else _meta_value(meta, "item_name", default="")
    if stock:
        lot = format_lot_for_masan_excel(
            lot_number_from=stock.lot_number_from,
            lot_number_to=stock.lot_number_to,
        )
    else:
        lot = format_lot_for_masan_excel(
            lot_number=str(_meta_value(meta, "lot", default="") or "") or None,
        )

    warehouse_name = order.warehouse.name if order.warehouse else ""
    storage_location = (
        detail.from_location.location_name
        if detail.from_location
        else _meta_value(meta, "storage_location")
    )
    locator = detail.to_location.location_name if detail.to_location else ""

    return [
        _format_inbound_datetime(order, meta),
        _meta_value(meta, "vehicle_no"),
        _meta_value(meta, "from_warehouse"),
        _meta_value(meta, "to_warehouse", default=warehouse_name),
        _meta_value(meta, "delivery"),
        _meta_value(meta, "nvt"),
        sku,
        item_name,
        lot,
        _meta_value(meta, "lot_status", default="GOOD"),
        allocation.quantity,
        _meta_value(meta, "pallet_count"),
        storage_location,
        locator,
        detail.device or None,
        _format_inbound_report_datetime(detail.start_time),
        _format_inbound_report_datetime(detail.end_time),
    ]


def _build_export_workbook(rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for col_idx, header in enumerate(INBOUND_REPORT_HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)

    last_data_row = DATA_START_ROW - 1
    for row_idx, row_values in enumerate(rows, start=DATA_START_ROW):
        last_data_row = row_idx
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(
                row=row_idx,
                column=col_idx,
                value=normalize_export_cell_value(col_idx, value),
            )

    style_masan_export_worksheet(ws, last_data_row=last_data_row)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _collect_inbound_export_rows(
    db: Session, order: InboundOrder
) -> list[list[Any]]:
    details = (
        db.query(InboundOrderDetail)
        .options(
            selectinload(InboundOrderDetail.allocations)
            .joinedload(InboundOrderAllocation.item_stock)
            .joinedload(ItemStock.item),
            joinedload(InboundOrderDetail.to_location),
            joinedload(InboundOrderDetail.from_location),
        )
        .filter(
            InboundOrderDetail.inbound_order_id == order.id,
            InboundOrderDetail.status == "completed",
        )
        .order_by(InboundOrderDetail.id)
        .all()
    )
    export_rows: list[list[Any]] = []
    for detail in details:
        allocations = sorted(detail.allocations, key=lambda a: a.id)
        if not allocations:
            continue
        for allocation in allocations:
            export_rows.append(_build_export_row(order, detail, allocation))
    return export_rows


def export_inbound_order_masan(db: Session, order_id: int) -> tuple[bytes, str]:
    order = (
        db.query(InboundOrder)
        .options(joinedload(InboundOrder.warehouse))
        .filter(InboundOrder.id == order_id)
        .first()
    )
    if not order:
        raise ValueError(f"Inbound order {order_id} not found")

    export_rows = _collect_inbound_export_rows(db, order)
    if not export_rows:
        raise ValueError(
            "Inbound order has no completed allocations to export"
        )

    content = _build_export_workbook(export_rows)
    filename = f"{order.order_code}_BaoCaoNhap.xlsx"
    return content, filename


_BULK_EXPORT_MAX_ORDERS = 100


def export_inbound_orders_masan(
    db: Session, order_ids: list[int]
) -> tuple[bytes, str]:
    unique_ids = list(dict.fromkeys(order_ids))
    if not unique_ids:
        raise ValueError("order_ids must not be empty")
    if len(unique_ids) > _BULK_EXPORT_MAX_ORDERS:
        raise ValueError(
            f"At most {_BULK_EXPORT_MAX_ORDERS} orders per export request"
        )

    orders = (
        db.query(InboundOrder)
        .options(joinedload(InboundOrder.warehouse))
        .filter(InboundOrder.id.in_(unique_ids))
        .all()
    )
    order_by_id = {o.id: o for o in orders}
    missing = [oid for oid in unique_ids if oid not in order_by_id]
    if missing:
        raise ValueError(
            f"Inbound orders not found: {sorted(missing)}"
        )

    export_rows: list[list[Any]] = []
    for oid in unique_ids:
        export_rows.extend(_collect_inbound_export_rows(db, order_by_id[oid]))

    if not export_rows:
        raise ValueError(
            "No completed inbound allocations to export for the selected orders"
        )

    content = _build_export_workbook(export_rows)
    first = order_by_id[unique_ids[0]]
    warehouse_label = (
        first.warehouse.code if first.warehouse and first.warehouse.code else str(first.warehouse_id)
    )
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    filename = f"BaoCaoNhap_{warehouse_label}_{len(unique_ids)}don_{stamp}.xlsx"
    return content, filename


_DAILY_CAPACITY_TEMPLATE = (
    Path(__file__).resolve().parents[2]
    / "static"
    / "templates"
    / "daily_capacity_template.xlsx"
)
_DAILY_CAPACITY_SUMMARY_HEADER_ROW = 2
_DAILY_CAPACITY_SUMMARY_START_ROW = 3
_DAILY_CAPACITY_SUMMARY_LAST_ROW = 5
_DAILY_CAPACITY_TABLE_HEADER_ROW = 6
_DAILY_CAPACITY_DATA_START_ROW = 7


def _daily_capacity_template_path() -> Path:
    if _DAILY_CAPACITY_TEMPLATE.is_file():
        return _DAILY_CAPACITY_TEMPLATE
    docs_dir = Path(__file__).resolve().parents[3] / "docs"
    for candidate in docs_dir.glob("*.xlsx"):
        if candidate.stat().st_size == 9601:
            return candidate
    raise ValueError("Daily capacity Excel template not found")


def _storage_locator_sql_sort_key():
    # Locator labels: G001..G420 (numeric suffix after single-letter prefix).
    return cast(func.substring(Location.location_name, 2), Integer)


def _excel_location_status(db_status: str | None) -> str:
    status = (db_status or "empty").strip()
    return "Trống" if status == "empty" else "Có hàng"


def _load_storage_stocks_by_location(
    db: Session, location_ids: list[int]
) -> dict[int, list[ItemStock]]:
    if not location_ids:
        return {}
    stocks = (
        db.query(ItemStock)
        .options(joinedload(ItemStock.item))
        .filter(
            ItemStock.location_id.in_(location_ids),
            ItemStock.is_active.is_(True),
            ItemStock.status == "available",
            ItemStock.available_quantity > 0,
            countable_stock_level_criterion(),
            positive_stock_quantity_criterion(),
        )
        .order_by(ItemStock.location_id.asc(), ItemStock.id.asc())
        .all()
    )
    by_location: dict[int, list[ItemStock]] = {}
    for stock in stocks:
        if stock.location_id is None:
            continue
        by_location.setdefault(stock.location_id, []).append(stock)
    return by_location


def _format_stock_inbound_date(created_at) -> str:
    if not created_at:
        return ""
    return created_at.strftime("%d/%m/%Y")


def _format_locator_stock_cells(
    stocks: list[ItemStock],
) -> tuple[Optional[str], Optional[str], Optional[str], Optional[str], Optional[str]]:
    if not stocks:
        return None, None, None, None, None
    sku_parts: list[str] = []
    name_parts: list[str] = []
    qty_parts: list[str] = []
    lot_parts: list[str] = []
    inbound_date_parts: list[str] = []
    for stock in stocks:
        item = stock.item
        sku_parts.append(item.sku if item else "")
        name_parts.append(item.name if item else "")
        qty_parts.append(str(int(stock.available_quantity)))
        lot_from = normalize_stored_lot_part(stock.lot_number_from)
        lot_to = normalize_stored_lot_part(stock.lot_number_to)
        lot_parts.append(
            format_lot_for_masan_excel(
                lot_number_from=lot_from,
                lot_number_to=lot_to,
            )
            or ""
        )
        inbound_date_parts.append(_format_stock_inbound_date(stock.created_at))
    sep = "\n"
    return (
        sep.join(sku_parts),
        sep.join(name_parts),
        sep.join(qty_parts),
        sep.join(lot_parts),
        sep.join(inbound_date_parts),
    )


def _list_storage_locations_for_daily_capacity(
    db: Session, warehouse_id: int
) -> list[tuple[Location, str]]:
    zone_codes = [c.strip() for c in settings.zone_storage if c and str(c).strip()]
    if not zone_codes:
        raise ValueError("zone_storage is not configured")

    locations = (
        db.query(Location)
        .join(Zone, Zone.id == Location.zone_id)
        .filter(
            Location.warehouse_id == warehouse_id,
            Location.is_active.is_(True),
            Zone.code.in_(zone_codes),
        )
        .order_by(
            _storage_locator_sql_sort_key().asc().nulls_last(),
            Location.location_name.asc(),
            Location.id.asc(),
        )
        .all()
    )
    if not locations:
        raise ValueError(
            "No active locations in storage zones for this warehouse"
        )

    assigned_ids = location_ids_with_assigned_qr()
    rows: list[tuple[Location, str]] = []
    for loc in locations:
        status = overlay_status_with_assign_cache(loc.status, loc.id, assigned_ids)
        rows.append((loc, status))
    return rows


def _build_daily_capacity_workbook(
    location_rows: list[tuple[Location, str]],
    stock_by_location: dict[int, list[ItemStock]],
) -> bytes:
    template_path = _daily_capacity_template_path()
    wb = load_masan_workbook(template_path)
    ws = wb.active

    total = len(location_rows)
    empty_count = sum(1 for _, status in location_rows if status == "empty")
    occupied_count = total - empty_count
    occupancy_pct = round((occupied_count / total) * 100, 2) if total else 0.0

    title = ws.cell(row=1, column=1).value
    if not title:
        title = "BIỂU MẪU CAPACITY CUỐI NGÀY"
    ws.cell(
        row=1,
        column=1,
        value=f"{title}\n(Ngày xuất: {date.today().strftime('%d/%m/%Y')})",
    )

    last_template_row = ws.max_row
    for row_idx in range(_DAILY_CAPACITY_SUMMARY_HEADER_ROW, last_template_row + 1):
        for col_idx in range(1, 4):
            ws.cell(row=row_idx, column=col_idx).value = None

    ws.cell(row=_DAILY_CAPACITY_SUMMARY_HEADER_ROW, column=1, value="Chỉ số")
    ws.cell(row=_DAILY_CAPACITY_SUMMARY_HEADER_ROW, column=2, value="Giá trị")

    ws.cell(row=_DAILY_CAPACITY_SUMMARY_START_ROW, column=1, value="Số lượng trống")
    ws.cell(row=_DAILY_CAPACITY_SUMMARY_START_ROW, column=2, value=empty_count)
    ws.cell(row=_DAILY_CAPACITY_SUMMARY_START_ROW + 1, column=1, value="Số lượng có hàng")
    ws.cell(
        row=_DAILY_CAPACITY_SUMMARY_START_ROW + 1,
        column=2,
        value=occupied_count,
    )
    ws.cell(
        row=_DAILY_CAPACITY_SUMMARY_LAST_ROW,
        column=1,
        value="Tỷ lệ % tồn kho",
    )
    ws.cell(
        row=_DAILY_CAPACITY_SUMMARY_LAST_ROW,
        column=2,
        value=f"{occupancy_pct}%",
    )

    ws.cell(row=_DAILY_CAPACITY_TABLE_HEADER_ROW, column=1, value="Locator")
    ws.cell(row=_DAILY_CAPACITY_TABLE_HEADER_ROW, column=2, value="Trạng thái")
    ws.cell(row=_DAILY_CAPACITY_TABLE_HEADER_ROW, column=3, value="Mã hàng")
    ws.cell(row=_DAILY_CAPACITY_TABLE_HEADER_ROW, column=4, value="Tên hàng")
    ws.cell(row=_DAILY_CAPACITY_TABLE_HEADER_ROW, column=5, value="Số lượng")
    ws.cell(row=_DAILY_CAPACITY_TABLE_HEADER_ROW, column=6, value="LOT")
    ws.cell(row=_DAILY_CAPACITY_TABLE_HEADER_ROW, column=7, value="Ngày nhập kho")

    last_data_row = _DAILY_CAPACITY_DATA_START_ROW - 1
    for offset, (loc, status) in enumerate(location_rows):
        row_idx = _DAILY_CAPACITY_DATA_START_ROW + offset
        last_data_row = row_idx
        ws.cell(row=row_idx, column=1, value=loc.location_name or "")
        ws.cell(row=row_idx, column=2, value=_excel_location_status(status))
        sku_cell, name_cell, qty_cell, lot_cell, inbound_date_cell = (
            _format_locator_stock_cells(stock_by_location.get(loc.id, []))
        )
        if sku_cell:
            ws.cell(row=row_idx, column=3, value=sku_cell)
            ws.cell(row=row_idx, column=4, value=name_cell)
            ws.cell(row=row_idx, column=5, value=qty_cell)
            if lot_cell:
                ws.cell(row=row_idx, column=6, value=str(lot_cell))
            if inbound_date_cell:
                ws.cell(row=row_idx, column=7, value=inbound_date_cell)

    style_daily_capacity_worksheet(
        ws,
        summary_header_row=_DAILY_CAPACITY_SUMMARY_HEADER_ROW,
        summary_last_row=_DAILY_CAPACITY_SUMMARY_LAST_ROW,
        table_header_row=_DAILY_CAPACITY_TABLE_HEADER_ROW,
        last_data_row=last_data_row,
    )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_daily_capacity_excel(db: Session, warehouse_id: int) -> tuple[bytes, str]:
    _ensure_warehouse_exists(db, warehouse_id)
    warehouse = (
        db.query(Warehouse).filter(Warehouse.id == warehouse_id).first()
    )
    location_rows = _list_storage_locations_for_daily_capacity(db, warehouse_id)
    location_ids = [loc.id for loc, _ in location_rows]
    stock_by_location = _load_storage_stocks_by_location(db, location_ids)
    content = _build_daily_capacity_workbook(location_rows, stock_by_location)
    warehouse_label = (warehouse.code if warehouse and warehouse.code else None) or str(
        warehouse_id
    )
    filename = f"Capacity_{warehouse_label}_{date.today().isoformat()}.xlsx"
    return content, filename


def _validate_masan_detail_filters(
    vehicle_no: Optional[str],
    item_id: Optional[int],
) -> None:
    has_vehicle = vehicle_no is not None and vehicle_no.strip() != ""
    has_item = item_id is not None

    if has_vehicle and has_item:
        raise ValueError(
            "Chỉ được truyền vehicle_no hoặc item_id, không truyền cả hai"
        )
    if not has_vehicle and not has_item:
        raise ValueError("Cần truyền vehicle_no hoặc item_id")


def _detail_query_options():
    return (
        selectinload(InboundOrderDetail.allocations)
        .joinedload(InboundOrderAllocation.item_stock)
        .joinedload(ItemStock.item),
        selectinload(InboundOrderDetail.allocations).joinedload(
            InboundOrderAllocation.unit
        ),
        joinedload(InboundOrderDetail.from_location),
        joinedload(InboundOrderDetail.to_location),
    )


def get_masan_inbound_order_details(
    db: Session,
    inbound_order_id: int,
    *,
    vehicle_no: Optional[str] = None,
    item_id: Optional[int] = None,
) -> list[InboundOrderDetailResponse]:
    _validate_masan_detail_filters(vehicle_no, item_id)

    order = (
        db.query(InboundOrder)
        .filter(InboundOrder.id == inbound_order_id)
        .first()
    )
    if not order:
        raise ValueError(f"Inbound order {inbound_order_id} not found")

    allocation_load, unit_load, from_loc, to_loc = _detail_query_options()
    query = (
        db.query(InboundOrderDetail)
        .options(allocation_load, unit_load, from_loc, to_loc)
        .filter(InboundOrderDetail.inbound_order_id == inbound_order_id)
    )

    if vehicle_no is not None and vehicle_no.strip():
        query = query.filter(
            InboundOrderDetail.details.op("->>")("vehicle_no") == vehicle_no.strip()
        )
    else:
        query = (
            query.join(InboundOrderAllocation)
            .join(ItemStock)
            .filter(ItemStock.item_id == item_id)
            .distinct()
        )

    details = query.order_by(InboundOrderDetail.id).all()
    return [_build_detail_response(d) for d in details]


def _location_name_sort_key(location_name: str | None) -> tuple[int, str]:
    name = (location_name or "").strip()
    digits = "".join(ch for ch in name if ch.isdigit())
    return (int(digits) if digits else -1, name)


def caller_masan_inbound_order(
    db: Session, location_ids: list[int], *, assign_robot_id: str | None = None
) -> dict[str, Any]:
    DEFAULT_GROUP = [0, 3, 2, 1]
    list_details: list[InboundOrderDetail] = []
    group = {}
    for location_id in location_ids:
        detail = (
            db.query(InboundOrderDetail)
            .filter(
                InboundOrderDetail.from_location_id == location_id,
                InboundOrderDetail.status == "initialize",
            )
            .order_by(desc(InboundOrderDetail.id))
            .first()
        )
        if detail:
            name = (
                db.query(Location.location_name)
                .filter(Location.id == location_id)
                .scalar()
            )
            warehouse_id = detail.inbound_order.warehouse_id
            name = (name or "").strip()
            m = re.search(r"(\d{2})$", name)
            pos = m.group(1) if m else ""
            if not pos:
                continue  
            point = int(pos) % 4
            group.setdefault(point, []).append(detail.id)

    for point in DEFAULT_GROUP:
        list_details.extend(group.get(point, []))

    if list_details:
        result = _run_inbound_task(db, warehouse_id, list_details, assign_robot_id=assign_robot_id)
        return result
    # detail_ids = list_details[:2]
    # list_details = list_details[2:]
    # cache_data = {
    #     "assign_robot_id": assign_robot_id,
    #     "list_details": list_details,
    # }
    # cache_set(f"inbound:caller_masan_inbound_order:{warehouse_id}", cache_data, -1)

    # job_ids = [
    #     accept_inbound_task_task.apply_async(
    #         kwargs={
    #             "detail_id": detail_id,
    #             **({"assign_robot_id": assign_robot_id} if assign_robot_id else {}),
    #         },
    #     ).id
    #     for detail_id in detail_ids
    # ]
    # return {
    #     "queued": len(detail_ids),
    #     "detail_ids": detail_ids,
    #     "job_ids": job_ids,
    # }

def _run_inbound_task(db: Session, warehouse_id: int, list_details: list[int], *, assign_robot_id: str | None = None, number_sending: int = 2) -> None:
    detail_ids = list_details[:number_sending]
    list_details = list_details[number_sending:]
    cache_data = {
        "assign_robot_id": assign_robot_id,
        "list_details": list_details,
    }
    cache_set(f"inbound:caller_masan_inbound_order:{warehouse_id}", cache_data, -1)

    job_ids = [
        accept_inbound_task_task.apply_async(
            kwargs={
                "detail_id": detail_id,
                **({"assign_robot_id": assign_robot_id} if assign_robot_id else {}),
            },
        ).id
        for detail_id in detail_ids
    ]
    return {
        "queued": len(detail_ids),
        "detail_ids": detail_ids,
        "job_ids": job_ids,
    }



def _product_fields_from_detail(
    detail: InboundOrderDetail,
) -> tuple[str | None, str | None, str | None]:
    extra = detail.details if isinstance(detail.details, dict) else {}
    sku = (extra.get("sku") or extra.get("part_number") or None)
    if sku is not None:
        sku = str(sku).strip() or None
    name = extra.get("item_name") or extra.get("product_name")
    if name is not None:
        name = str(name).strip() or None
    lot: str | None = None

    for allocation in detail.allocations:
        stock = allocation.item_stock
        item = stock.item if stock else None
        if item:
            sku = sku or (item.sku or None)
            name = name or (item.name or None)
        if stock and not lot:
            lot = format_lot_number_display(
                stock.lot_number_from,
                stock.lot_number_to,
            )

    return sku, name, lot


def list_inbound_details_awaiting_robot(
    db: Session,
    order_ids: list[int],
) -> InboundAssignedDetailsResponse:
    unique_order_ids = list(dict.fromkeys(order_ids))
    if not unique_order_ids:
        return InboundAssignedDetailsResponse(order_ids=[], details=[])

    existing_ids = {
        row[0]
        for row in db.query(InboundOrder.id)
        .filter(InboundOrder.id.in_(unique_order_ids))
        .all()
    }
    missing = [oid for oid in unique_order_ids if oid not in existing_ids]
    if missing:
        raise ValueError(f"Inbound order(s) not found: {', '.join(map(str, missing))}")

    allocation_load, unit_load, from_loc, to_loc = _detail_query_options()
    details = (
        db.query(InboundOrderDetail)
        .options(allocation_load, unit_load, from_loc, to_loc)
        .filter(
            InboundOrderDetail.inbound_order_id.in_(unique_order_ids),
            InboundOrderDetail.status == "initialize",
            InboundOrderDetail.detail_type != "manual",
            InboundOrderDetail.from_location_id.isnot(None),
            InboundOrderDetail.to_location_id.isnot(None),
        )
        .order_by(InboundOrderDetail.inbound_order_id, InboundOrderDetail.id)
        .all()
    )

    items: list[InboundAssignedDetailItem] = []
    for detail in details:
        sku, name, lot = _product_fields_from_detail(detail)
        items.append(
            InboundAssignedDetailItem(
                detail_id=detail.id,
                inbound_order_id=detail.inbound_order_id,
                product_sku=sku,
                product_name=name,
                lot_number=lot,
                from_location=location_display_label(detail.from_location),
                to_location=location_display_label(detail.to_location),
                from_location_code=(
                    detail.from_location.location_code
                    if detail.from_location
                    else None
                ),
                status=detail.status,
            )
        )

    return InboundAssignedDetailsResponse(
        order_ids=unique_order_ids,
        details=items,
    )


def clear_zone_inbound(db: Session, warehouse_id: int) -> int:
    list_zones = settings.zone_inbound
    zone_filter = or_(
        Zone.name.in_(list_zones),
        Zone.code.in_(list_zones),
    )

    logger.info(f"Clearing zones {list_zones} for warehouse {warehouse_id}")
    stocks = (
        db.query(ItemStock)
        .join(Location, Location.id == ItemStock.location_id)
        .join(Zone, Zone.id == Location.zone_id)
        .filter(
            Location.warehouse_id == warehouse_id,
            zone_filter,
            ItemStock.is_active.is_(True),
        )
        .all()
    )

    item_stock_ids = [stock.id for stock in stocks]

    if not item_stock_ids:
        details = []
    else:
        details = (
            db.query(InboundOrderDetail)
            .join(
                InboundOrderAllocation,
                InboundOrderAllocation.inbound_order_detail_id == InboundOrderDetail.id,
            )
            .filter(InboundOrderAllocation.item_stock_id.in_(item_stock_ids))
            .distinct()
            .all()
        )

    logger.info(f"Found {len(stocks)} stocks to clear")

    for stock in stocks:
        stock.is_active = False

    for detail in details:
        detail.status = "not call"
    db.commit()
    return len(stocks)
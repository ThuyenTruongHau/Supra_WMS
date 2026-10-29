from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from typing import Any, Optional
from sqlalchemy import func, Integer, cast, or_
import uuid
import json

from app.core.config import settings
from openpyxl import Workbook
from datetime import datetime

from app.modules.masan.masan_excel_io import load_masan_workbook
from app.modules.masan.masan_oubound_auto_service import _exist_robot_task
from sqlalchemy.orm import Session, joinedload, selectinload

from app.core.cache import cache_set, cache_get, cache_scan_keys, cache_delete, cache_delete_pattern, cache_replace_if_unchanged
from app.modules.masan.masan_outbound_excel import (
    COLUMN_ALIASES,
    DATA_START_ROW,
    HEADER_ROW,
    OUTBOUND_REPORT_HEADERS,
    TRIP_PLACEHOLDERS,
    normalize_outbound_report_cell_value,
    style_outbound_report_worksheet,
)
from app.modules.robot.robot_model import RobotTask
from app.modules.warehouse.outbound_order.outbound_order_schema import (
    AllocationOutboundTaskExecute,
    CalculateOutboundDetail,
    DetailForCalculate,
    OutboundRobotTaskCreate,
)
from app.modules.warehouse.outbound_order.outbound_order_service import (
    calculate_outbound_order,
    execute_outbound_task,
    greedy_allocate_stocks_to_lines,
    lacked_details,
    _settle_outbound_stock,
)
from app.modules.masan.masan_outbound_so_excel import (
    build_so_customer_sheet,
    sanitize_sheet_name,
)
from app.modules.warehouse.outbound_order.outbound_order_model import (
    OutboundOrder,
    OutboundOrderAllocation,
    OutboundOrderDetail,
)
from app.modules.masan.masan_schema import (
    MasanOutboundLineItem,
    MasanOutboundParseResponse,
    MasanOutboundPreviewRow,
)
from app.modules.warehouse.location_map.location_model import Location
from app.modules.warehouse.item.item_model import Item
from app.modules.warehouse.item_stock.item_stock_model import (
    ItemStock,
    countable_stock_level_criterion,
)
from app.modules.warehouse.lot_number_utils import (
    coerce_masan_lot_cell,
    format_lot_for_masan_excel,
    format_lot_value_for_display,
    normalize_masan_lot_cell,
    parse_legacy_lot_number,
)
from app.modules.warehouse.warehouse_zone.warehouse_model import Zone
from app.modules.warehouse.warehouse_zone.warehouse_service import _ensure_warehouse_exists
from app.core.logger import get_logger
from app.socket.ws_events import (
    EVENT_MASAN_SORTING_STOCK_READY,
    publish_masan_sorting_zone,
)

logger = get_logger("main")


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower()


def _build_column_map(header_row: list[Any]) -> dict[str, int]:
    col_map: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        name = _normalize_header(cell)
        for field, aliases in COLUMN_ALIASES.items():
            if name in aliases:
                col_map[field] = idx
    return col_map


def _direct(row: list[Any], col_map: dict[str, int], field: str) -> Optional[str]:
    idx = col_map.get(field)
    if idx is None or idx >= len(row) or row[idx] in (None, ""):
        return None
    return str(row[idx]).strip()


def _is_blank_cell(value: Any) -> bool:
    return value is None or str(value).strip() == ""


def _build_merge_anchor_map(ws) -> dict[tuple[int, int], tuple[int, int]]:
    """Map each (row, col) in a merge range to its top-left cell (1-based)."""
    anchors: dict[tuple[int, int], tuple[int, int]] = {}
    for merged in ws.merged_cells.ranges:
        top_left = (merged.min_row, merged.min_col)
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                anchors[(row, col)] = top_left
    return anchors


def _read_merged_row(
    ws,
    row_idx: int,
    num_cols: int,
    merge_anchors: dict[tuple[int, int], tuple[int, int]],
) -> list[Any]:
    row: list[Any] = []
    for col in range(1, num_cols + 1):
        anchor = merge_anchors.get((row_idx, col))
        if anchor:
            row.append(ws.cell(anchor[0], anchor[1]).value)
        else:
            row.append(ws.cell(row_idx, col).value)
    return row


def _merged_field_raw(
    row: list[Any], col_map: dict[str, int], field: str
) -> Any:
    idx = col_map.get(field)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _maybe_reset_item_carry_on_new_vehicle(
    row: list[Any],
    col_map: dict[str, int],
    carry: dict[str, Any],
) -> None:
    idx = col_map.get("vehicle_no")
    if idx is None or idx >= len(row) or _is_blank_cell(row[idx]):
        return
    vehicle = str(row[idx]).strip()
    previous = carry.get("_vehicle_anchor")
    if previous is not None and previous != vehicle:
        for key in (
            "sku",
            "item_name",
            "lot",
            "lot_status",
            "quantity",
            "pallet_count",
        ):
            carry.pop(key, None)
    carry["_vehicle_anchor"] = vehicle


def _cell_raw(
    row: list[Any],
    col_map: dict[str, int],
    field: str,
    carry: dict[str, Any],
) -> Any:
    idx = col_map.get(field)
    if idx is None:
        return carry.get(field)

    raw = row[idx] if idx < len(row) else None
    if _is_blank_cell(raw):
        return carry.get(field)

    carry[field] = raw
    return raw


def _normalize_trip(value: Optional[str], carry: Optional[str]) -> Optional[str]:
    if value is None:
        return carry
    cleaned = value.strip()
    if not cleaned or cleaned in TRIP_PLACEHOLDERS:
        return carry
    return cleaned


def _cell_value(
    row: list[Any],
    col_map: dict[str, int],
    field: str,
    carry: dict[str, Any],
) -> Optional[str]:
    raw = _cell_raw(row, col_map, field, carry)
    if _is_blank_cell(raw):
        prev = carry.get(field)
        if prev is None or _is_blank_cell(prev):
            return None
        return str(prev).strip()

    value = str(raw).strip()
    if field == "trip":
        normalized = _normalize_trip(value, str(carry.get(field) or "").strip() or None)
        if normalized:
            carry[field] = normalized
        return normalized

    carry[field] = value
    return value


def _parse_quantity(raw: Any) -> tuple[int, Optional[str]]:
    if raw in (None, ""):
        return 0, "Số lượng trống"
    try:
        qty = int(float(str(raw).replace(",", "")))
    except (TypeError, ValueError):
        return 0, f"Số lượng không hợp lệ: {raw!r}"
    if qty <= 0:
        return qty, "Số lượng phải lớn hơn 0"
    return qty, None


def _validate_lot(raw: Optional[str], outbound_type: str) -> Optional[str]:
    if raw is None or str(raw).strip() == "":
        return None
    lot = coerce_masan_lot_cell(str(raw))
    if lot is None or str(lot).strip() == "":
        return None
    if outbound_type == "auto":
        single = normalize_masan_lot_cell(lot)
        if not single:
            try:
                parse_legacy_lot_number(str(lot))
            except ValueError as exc:
                return str(exc)
            return (
                "LOT Masan (auto) phải là một mốc DDMMYY 6 chữ số, "
                "không dùng khoảng lô"
            )
    return None


def parse_masan_outbound_rows(content: bytes) -> list[dict[str, Any]]:
    wb = load_masan_workbook(content, data_only=True)
    ws = wb.active
    merge_anchors = _build_merge_anchor_map(ws)
    num_cols = ws.max_column

    header_row = _read_merged_row(ws, HEADER_ROW, num_cols, merge_anchors)
    col_map = _build_column_map(header_row)

    if "sku" not in col_map:
        raise ValueError('Không tìm thấy cột "Mã Item" ở hàng 2')

    rows: list[dict[str, Any]] = []
    carry: dict[str, Any] = {}

    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        row = _read_merged_row(ws, row_idx, num_cols, merge_anchors)
        _maybe_reset_item_carry_on_new_vehicle(row, col_map, carry)

        merged_sku = _merged_field_raw(row, col_map, "sku")
        merged_locator = _merged_field_raw(row, col_map, "locator")
        merged_qty = _merged_field_raw(row, col_map, "quantity")
        has_signal = not (
            _is_blank_cell(merged_sku)
            and _is_blank_cell(merged_locator)
            and _is_blank_cell(merged_qty)
        )
        if not has_signal:
            continue

        sku_raw = _cell_raw(row, col_map, "sku", carry)
        sku = str(sku_raw).strip() if not _is_blank_cell(sku_raw) else ""
        if not sku:
            continue

        quantity_raw = _cell_raw(row, col_map, "quantity", carry)
        quantity, qty_error = _parse_quantity(quantity_raw)

        lot_raw = _cell_value(row, col_map, "lot", carry)
        lot = coerce_masan_lot_cell(lot_raw) if lot_raw else None

        rows.append(
            {
                "row_no": row_idx,
                "vehicle_no": _cell_value(row, col_map, "vehicle_no", carry),
                "customer_name": _cell_value(row, col_map, "customer_name", carry),
                "trip": _cell_value(row, col_map, "trip", carry),
                "nvt": _cell_value(row, col_map, "nvt", carry),
                "sku": sku,
                "item_name": _cell_value(row, col_map, "item_name", carry),
                "lot": lot,
                "lot_status": _cell_value(row, col_map, "lot_status", carry),
                "quantity": quantity,
                "quantity_error": qty_error,
                "pallet_count": _cell_value(row, col_map, "pallet_count", carry),
                "locator": _direct(row, col_map, "locator"),
            }
        )

    if not rows:
        raise ValueError("Không có dòng hợp lệ trong file Excel")

    return rows


def _lookup_item(db: Session, warehouse_id: int, sku: str) -> Optional[Item]:
    return (
        db.query(Item)
        .filter(
            Item.warehouse_id == warehouse_id,
            Item.sku == sku,
            Item.is_active.is_(True),
        )
        .first()
    )


def _validate_outbound_import_locator(
    db: Session,
    *,
    warehouse_id: int,
    row_no: int,
    item: Item,
    quantity: int,
    locator: str,
) -> Optional[str]:
    """Return error message when BM.04 row has Locator but stock cannot satisfy the line."""
    locator_str = str(locator).strip()
    if not locator_str:
        return None

    location = (
        db.query(Location)
        .filter(
            Location.warehouse_id == warehouse_id,
            Location.is_active.is_(True),
            or_(
                Location.location_name == locator_str,
                Location.location_code == locator_str,
            ),
        )
        .first()
    )
    if not location:
        return (
            f"Dòng {row_no}: Không tìm thấy Locator \"{locator_str}\" trong kho"
        )

    stocks = (
        db.query(ItemStock)
        .filter(
            ItemStock.location_id == location.id,
            ItemStock.is_active.is_(True),
            ItemStock.status.in_(["available", "split"]),
            ItemStock.quantity > 0,
        )
        .all()
    )
    if not stocks:
        return f"Dòng {row_no}: Locator \"{locator_str}\" không có tồn kho"

    item_stocks = [s for s in stocks if s.item_id == item.id]
    if not item_stocks:
        return (
            f"Dòng {row_no}: Locator \"{locator_str}\" không có mã Item "
            f"\"{item.sku}\" (khác loại hàng tại vị trí này)"
        )

    available = sum(int(s.available_quantity or 0) for s in item_stocks)
    if available < quantity:
        return (
            f"Dòng {row_no}: Locator \"{locator_str}\" chỉ còn {available} "
            f"khả dụng, yêu cầu {quantity}"
        )

    return None


def _build_line_item_details(row: dict[str, Any]) -> dict[str, Any]:
    details: dict[str, Any] = {}
    for key in (
        "vehicle_no",
        "customer_name",
        "trip",
        "nvt",
        "lot_status",
        "pallet_count",
        "locator",
    ):
        value = row.get(key)
        if value is not None and str(value).strip() != "":
            details[key] = value
    lot = coerce_masan_lot_cell(row.get("lot"))
    if lot is not None and str(lot).strip() != "":
        lot_text = str(lot).strip()
        details["lot_number"] = (
            format_lot_value_for_display(lot_text) or lot_text
        )
    return details


def parse_masan_outbound_preview(
    db: Session,
    *,
    warehouse_id: int,
    content: bytes,
    outbound_type: str,
) -> MasanOutboundParseResponse:
    _ensure_warehouse_exists(db, warehouse_id)

    parsed_rows = parse_masan_outbound_rows(content)
    sku_cache: dict[str, Optional[Item]] = {}

    preview_rows: list[MasanOutboundPreviewRow] = []
    line_items: list[MasanOutboundLineItem] = []
    warnings: list[str] = []
    invalid_count = 0
    locator_import_errors: list[str] = []

    for row in parsed_rows:
        sku = row["sku"]
        if sku not in sku_cache:
            sku_cache[sku] = _lookup_item(db, warehouse_id, sku)

        item = sku_cache[sku]
        lot_raw = row.get("lot")

        errors: list[str] = []
        if row.get("quantity_error"):
            errors.append(row["quantity_error"])
        if item is None:
            errors.append(f"Không tìm thấy sản phẩm SKU '{sku}' trong kho")
        lot_error = _validate_lot(lot_raw, outbound_type)
        if lot_error:
            errors.append(lot_error)

        locator_raw = row.get("locator")
        if (
            locator_raw is not None
            and str(locator_raw).strip() != ""
            and item is not None
            and not row.get("quantity_error")
        ):
            locator_error = _validate_outbound_import_locator(
                db,
                warehouse_id=warehouse_id,
                row_no=int(row["row_no"]),
                item=item,
                quantity=int(row["quantity"]),
                locator=str(locator_raw),
            )
            if locator_error:
                errors.append(locator_error)
                locator_import_errors.append(locator_error)

        preview_rows.append(
            MasanOutboundPreviewRow(
                row_no=row["row_no"],
                vehicle_no=row.get("vehicle_no"),
                customer_name=row.get("customer_name"),
                trip=row.get("trip"),
                nvt=row.get("nvt"),
                sku=sku,
                item_name=row.get("item_name"),
                lot_number=str(lot_raw).strip() if lot_raw not in (None, "") else None,
                lot_status=row.get("lot_status"),
                quantity=row.get("quantity") or 0,
                pallet_count=row.get("pallet_count"),
                locator=row.get("locator"),
                item_id=item.id if item else None,
                unit_id=item.base_unit_id if item else None,
                error="; ".join(errors) if errors else None,
            )
        )

        if errors or item is None:
            invalid_count += 1
            continue

        line_items.append(
            MasanOutboundLineItem(
                item_id=item.id,
                quantity=row["quantity"],
                unit_id=item.base_unit_id,
                detail_type=outbound_type,
                details=_build_line_item_details(row),
            )
        )

    import_errors: list[str] = list(locator_import_errors)
    can_import = bool(line_items) and not import_errors

    if not line_items and not import_errors:
        import_errors.append(
            "Không có dòng hợp lệ để tạo đơn xuất. "
            "Kiểm tra SKU, số lượng và LOT trong file."
        )

    if invalid_count:
        warnings.append(f"{invalid_count} dòng bị bỏ qua khỏi payload tạo đơn")

    if import_errors:
        warnings.insert(
            0,
            f"Không thể import: {len(import_errors)} lỗi (xem import_errors / error từng dòng).",
        )

    return MasanOutboundParseResponse(
        preview_rows=preview_rows,
        line_items=line_items if can_import else [],
        total_rows=len(preview_rows),
        valid_rows=len(preview_rows) - invalid_count,
        invalid_rows=invalid_count,
        warnings=warnings,
        import_errors=import_errors,
        can_import=can_import,
    )


def _unique_keep_order(values: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        cleaned = value.strip()
        if not cleaned or cleaned in seen:
            continue
        seen.add(cleaned)
        result.append(cleaned)
    return result


def _resolve_locator(detail: OutboundOrderDetail) -> str:
    outbound_allocs = [
        alloc
        for alloc in detail.allocations
        if alloc.allocation_type == "outbound"
    ]
    if outbound_allocs:
        locations: list[str] = []
        seen: set[str] = set()
        for alloc in sorted(outbound_allocs, key=lambda a: a.id):
            loc = alloc.from_location
            label = (
                loc.location_name or loc.location_code
                if loc
                else None
            )
            if label and label not in seen:
                seen.add(label)
                locations.append(label)
        if locations:
            return ", ".join(locations)

    meta = detail.details or {}
    locator = meta.get("locator")
    if locator and str(locator).strip():
        return str(locator).strip()

    return "..."


def _customer_group_key(detail: OutboundOrderDetail) -> str:
    meta = detail.details or {}
    name = str(meta.get("customer_name") or "").strip()
    return name or "Unknown"


def _build_so_data_row(detail: OutboundOrderDetail, stt: int) -> list[Any]:
    meta = detail.details or {}
    lot_status = meta.get("lot_status")
    if not lot_status or str(lot_status).strip() == "":
        lot_status = "GOOD"
    pallet_count = meta.get("pallet_count")
    return [
        stt,
        detail.item.sku if detail.item else "",
        detail.item.name if detail.item else "",
        format_lot_for_masan_excel(
            lot_number=str(meta.get("lot_number") or "") or None,
        )
        or None,
        lot_status,
        detail.quantity,
        pallet_count if pallet_count not in (None, "") else None,
        _resolve_locator(detail),
    ]


def export_outbound_order_so(db: Session, order_id: int) -> tuple[bytes, str]:
    order = (
        db.query(OutboundOrder)
        .filter(OutboundOrder.id == order_id)
        .first()
    )
    if not order:
        raise ValueError(f"Outbound order {order_id} not found")

    details = (
        db.query(OutboundOrderDetail)
        .options(
            joinedload(OutboundOrderDetail.item),
            selectinload(OutboundOrderDetail.allocations).joinedload(
                OutboundOrderAllocation.from_location
            ),
        )
        .filter(OutboundOrderDetail.outbound_order_id == order.id)
        .order_by(OutboundOrderDetail.id)
        .all()
    )
    if not details:
        raise ValueError("Outbound order has no details to export")

    grouped: dict[str, list[OutboundOrderDetail]] = defaultdict(list)
    for detail in details:
        grouped[_customer_group_key(detail)].append(detail)

    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    used_sheet_names: set[str] = set()
    for customer_name in sorted(grouped.keys(), key=lambda name: (name == "Unknown", name)):
        customer_details = grouped[customer_name]
        sheet_name = sanitize_sheet_name(customer_name, used=used_sheet_names)
        ws = wb.create_sheet(title=sheet_name)

        trips: list[str] = []
        vehicles: list[str] = []
        nvts: list[str] = []
        for detail in customer_details:
            meta = detail.details or {}
            for key, bucket in (
                ("trip", trips),
                ("vehicle_no", vehicles),
                ("nvt", nvts),
            ):
                value = meta.get(key)
                if value is not None and str(value).strip() != "":
                    bucket.append(str(value).strip())

        data_rows = [
            _build_so_data_row(detail, stt)
            for stt, detail in enumerate(
                sorted(customer_details, key=lambda d: d.id),
                start=1,
            )
        ]
        build_so_customer_sheet(
            ws,
            customer_name=customer_name,
            trips=_unique_keep_order(trips),
            vehicles=_unique_keep_order(vehicles),
            nvts=_unique_keep_order(nvts),
            data_rows=data_rows,
        )

    buffer = BytesIO()
    wb.save(buffer)
    filename = f"{order.order_code}_LayHangSO.xlsx"
    return buffer.getvalue(), filename


def _meta_value(meta: dict[str, Any], key: str, default: Any = "") -> Any:
    value = meta.get(key)
    if value is None or value == "":
        return default
    return value


def _format_outbound_report_datetime(dt: datetime | None) -> str:
    if not dt:
        return ""
    return dt.strftime("%d/%m/%Y\n%H:%M:%S")


def _format_report_distance_m(distance_m: float | None) -> float | None:
    if distance_m is None:
        return None
    return round(float(distance_m), 3)


def _format_report_battery(battery: Any) -> str | None:
    if battery is None:
        return None
    return str(battery).strip() or None


def _location_display_label(loc: Any) -> str:
    if loc is None:
        return ""
    label = loc.location_name or loc.location_code
    if label is None:
        return ""
    cleaned = str(label).strip()
    return cleaned


def _resolve_detail_sorting_location(
    detail: OutboundOrderDetail,
    *,
    db: Session | None = None,
) -> str:
    """Sorting/CC label for the detail line — reused on every allocation row."""
    label = _location_display_label(detail.sorting_location)
    if label:
        return label
    if detail.sorting_location_id is not None and db is not None:
        loc = db.get(Location, detail.sorting_location_id)
        label = _location_display_label(loc)
        if label:
            return label
    meta = detail.details if isinstance(detail.details, dict) else {}
    for key in ("sorting_location", "cc_location", "vi_tri_chia"):
        raw = meta.get(key)
        if raw is not None and str(raw).strip():
            return str(raw).strip()
    if db is not None:
        wh_id = detail.outbound_order.warehouse_id if detail.outbound_order else None
        if wh_id is None and detail.outbound_order_id:
            wh_id = (
                db.query(OutboundOrder.warehouse_id)
                .filter(OutboundOrder.id == detail.outbound_order_id)
                .scalar()
            )
        if wh_id is not None:
            bucket = cache_get(
                f"outbound:assigned_cc_details:{wh_id}:{detail.id}"
            )
            if isinstance(bucket, dict) and bucket.get("location_id") is not None:
                loc = db.get(Location, int(bucket["location_id"]))
                label = _location_display_label(loc)
                if label:
                    return label
    return ""


def _resolve_allocation_locator(
    allocation: OutboundOrderAllocation,
    detail: OutboundOrderDetail,
) -> str:
    label = _location_display_label(allocation.from_location)
    if label:
        return label
    locator = _meta_value(meta=detail.details or {}, key="locator", default="")
    return str(locator).strip() if locator else ""


_OUTBOUND_REPORT_DETAIL_LOAD_OPTIONS = (
    joinedload(OutboundOrderDetail.item),
    joinedload(OutboundOrderDetail.outbound_order),
    joinedload(OutboundOrderDetail.sorting_location),
    selectinload(OutboundOrderDetail.allocations).joinedload(
        OutboundOrderAllocation.item_stock
    ),
    selectinload(OutboundOrderDetail.allocations).joinedload(
        OutboundOrderAllocation.from_location
    ),
    selectinload(OutboundOrderDetail.allocations).joinedload(
        OutboundOrderAllocation.to_location
    ),
)


def _build_outbound_report_row(
    detail: OutboundOrderDetail,
    allocation: OutboundOrderAllocation,
    *,
    meta: dict[str, Any] | None = None,
    item: Any | None = None,
    detail_sorting_location: str,
) -> list[Any]:
    meta = meta if meta is not None else (detail.details or {})
    stock: ItemStock | None = allocation.item_stock
    if item is None:
        item = detail.item

    if stock:
        lot = format_lot_for_masan_excel(
            lot_number_from=stock.lot_number_from,
            lot_number_to=stock.lot_number_to,
        )
    else:
        lot = format_lot_for_masan_excel(
            lot_number=str(_meta_value(meta, "lot_number", default="") or "") or None,
        )

    lot_status = _meta_value(meta, "lot_status", default="GOOD")
    if not str(lot_status).strip():
        lot_status = "GOOD"
    pallet_count = meta.get("pallet_count")

    return [
        _meta_value(meta, "vehicle_no"),
        _meta_value(meta, "customer_name"),
        _meta_value(meta, "trip"),
        _meta_value(meta, "nvt"),
        item.sku if item else "",
        item.name if item else "",
        lot,
        lot_status,
        allocation.quantity,
        pallet_count if pallet_count not in (None, "") else None,
        _resolve_allocation_locator(allocation, detail),
        (
            str(allocation.to_location.location_name).strip()
            if allocation.to_location and allocation.to_location.location_name
            else None
        ),
        (allocation.sorting_persons or "").strip() or None,
        _format_report_distance_m(allocation.distance),
        allocation.device or None,
        _format_outbound_report_datetime(allocation.start_time),
        _format_outbound_report_datetime(allocation.end_time),
        _format_report_battery(allocation.start_battery),
        _format_report_battery(allocation.end_battery),
        detail_sorting_location,
        _format_outbound_report_datetime(allocation.start_sorting_time),
        _format_outbound_report_datetime(allocation.end_sorting_time),
    ]


def _build_outbound_report_workbook(rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for col_idx, header in enumerate(OUTBOUND_REPORT_HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)

    last_data_row = DATA_START_ROW - 1
    for row_idx, row_values in enumerate(rows, start=DATA_START_ROW):
        last_data_row = row_idx
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(
                row=row_idx,
                column=col_idx,
                value=normalize_outbound_report_cell_value(col_idx, value),
            )

    style_outbound_report_worksheet(ws, last_data_row=last_data_row)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _fetch_outbound_report_details_by_order(
    db: Session,
    order_ids: list[int],
) -> dict[int, list[OutboundOrderDetail]]:
    if not order_ids:
        return {}
    details = (
        db.query(OutboundOrderDetail)
        .options(*_OUTBOUND_REPORT_DETAIL_LOAD_OPTIONS)
        .filter(
            OutboundOrderDetail.outbound_order_id.in_(order_ids),
            OutboundOrderDetail.status == "completed",
        )
        .order_by(
            OutboundOrderDetail.outbound_order_id,
            OutboundOrderDetail.id,
        )
        .all()
    )
    by_order: dict[int, list[OutboundOrderDetail]] = defaultdict(list)
    for detail in details:
        by_order[detail.outbound_order_id].append(detail)
    return dict(by_order)


def _iter_outbound_report_rows(
    details: list[OutboundOrderDetail],
    *,
    db: Session | None = None,
) -> list[list[Any]]:
    export_rows: list[list[Any]] = []
    for detail in details:
        outbound_allocs = [
            a for a in detail.allocations if a.allocation_type == "outbound"
        ]
        outbound_allocs.sort(key=lambda a: a.id)
        # Same detail sorting label on every allocation row for this line.
        detail_sorting_location = _resolve_detail_sorting_location(detail, db=db)
        meta = detail.details or {}
        item = detail.item
        for allocation in outbound_allocs:
            export_rows.append(
                _build_outbound_report_row(
                    detail,
                    allocation,
                    meta=meta,
                    item=item,
                    detail_sorting_location=detail_sorting_location,
                )
            )
    return export_rows


def _collect_outbound_report_rows(
    db: Session, order: OutboundOrder
) -> list[list[Any]]:
    by_order = _fetch_outbound_report_details_by_order(db, [order.id])
    return _iter_outbound_report_rows(by_order.get(order.id, []), db=db)


def export_outbound_order_report_masan(
    db: Session, order_id: int
) -> tuple[bytes, str]:
    order = (
        db.query(OutboundOrder)
        .filter(OutboundOrder.id == order_id)
        .first()
    )
    if not order:
        raise ValueError(f"Outbound order {order_id} not found")

    export_rows = _collect_outbound_report_rows(db, order)
    if not export_rows:
        raise ValueError(
            "Outbound order has no completed allocations to export"
        )

    content = _build_outbound_report_workbook(export_rows)
    filename = f"{order.order_code}_BaoCaoXuat.xlsx"
    return content, filename


_BULK_OUTBOUND_REPORT_MAX_ORDERS = 100


def export_outbound_orders_report_masan(
    db: Session, order_ids: list[int]
) -> tuple[bytes, str]:
    unique_ids = list(dict.fromkeys(order_ids))
    if not unique_ids:
        raise ValueError("order_ids must not be empty")
    if len(unique_ids) > _BULK_OUTBOUND_REPORT_MAX_ORDERS:
        raise ValueError(
            f"At most {_BULK_OUTBOUND_REPORT_MAX_ORDERS} orders per export request"
        )

    orders = (
        db.query(OutboundOrder)
        .filter(OutboundOrder.id.in_(unique_ids))
        .all()
    )
    order_by_id = {o.id: o for o in orders}
    missing = [oid for oid in unique_ids if oid not in order_by_id]
    if missing:
        raise ValueError(
            f"Outbound orders not found: {sorted(missing)}"
        )

    details_by_order = _fetch_outbound_report_details_by_order(db, unique_ids)
    export_rows: list[list[Any]] = []
    for oid in unique_ids:
        export_rows.extend(
            _iter_outbound_report_rows(details_by_order.get(oid, []), db=db)
        )

    if not export_rows:
        raise ValueError(
            "No completed outbound allocations to export for the selected orders"
        )

    content = _build_outbound_report_workbook(export_rows)
    filename = f"BaoCaoXuat_{len(unique_ids)}don.xlsx"
    return content, filename


def _resolve_line_sku(detail: OutboundOrderDetail, meta: dict[str, Any]) -> str | None:
    raw = meta.get("sku")
    if raw is not None and str(raw).strip() != "":
        return str(raw).strip()
    item = detail.item
    if item is not None and item.sku:
        return item.sku
    return None


def _sku_from_allocation(allocation: OutboundOrderAllocation) -> str | None:
    detail = allocation.outbound_order_detail
    if detail is not None:
        meta = detail.details if isinstance(detail.details, dict) else {}
        resolved = _resolve_line_sku(detail, meta)
        if resolved:
            return resolved
    stock = allocation.item_stock
    if stock is not None and stock.item is not None and stock.item.sku:
        return str(stock.item.sku).strip() or None
    return None


def _locator_name_for_allocation(allocation: OutboundOrderAllocation) -> str | None:
    loc = allocation.from_location
    if loc is not None:
        name = (loc.location_name or "").strip()
        if name:
            return name
        code = (loc.location_code or "").strip()
        if code:
            return code
    detail = allocation.outbound_order_detail
    if detail is not None and isinstance(detail.details, dict):
        raw = detail.details.get("locator")
        if raw is not None and str(raw).strip():
            return str(raw).strip()
    return None


def _enrich_stock_ready_payloads(
    db: Session,
    rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    if not rows:
        return rows
    alloc_ids: list[int] = []
    for row in rows:
        raw_id = row.get("id")
        if raw_id is not None:
            alloc_ids.append(int(raw_id))
    if not alloc_ids:
        return rows
    allocations = (
        db.query(OutboundOrderAllocation)
        .options(
            joinedload(OutboundOrderAllocation.from_location),
            joinedload(OutboundOrderAllocation.outbound_order_detail).joinedload(
                OutboundOrderDetail.item
            ),
            joinedload(OutboundOrderAllocation.item_stock).joinedload(ItemStock.item),
        )
        .filter(OutboundOrderAllocation.id.in_(alloc_ids))
        .all()
    )
    by_id = {a.id: a for a in allocations}
    enriched: list[dict[str, Any]] = []
    for row in rows:
        copy = dict(row)
        alloc_id = copy.get("id")
        if alloc_id is not None:
            allocation = by_id.get(int(alloc_id))
            if allocation is not None:
                if not (copy.get("sku") and str(copy.get("sku")).strip()):
                    sku = _sku_from_allocation(allocation)
                    if sku:
                        copy["sku"] = sku
                locator = _locator_name_for_allocation(allocation)
                if locator:
                    copy["locator"] = locator
                detail = allocation.outbound_order_detail
                if detail is not None:
                    cache_line = _detail_to_cache_line(detail)
                    for key in (
                        "outbound_order_id",
                        "item_id",
                        "unit",
                        "detail_type",
                        "vehicle_no",
                        "customer_name",
                        "trip",
                        "nvt",
                        "lot_number",
                        "pallet_count",
                    ):
                        existing = copy.get(key)
                        if existing is not None and str(existing).strip() != "":
                            continue
                        value = cache_line.get(key)
                        if value is not None and str(value).strip() != "":
                            copy[key] = value
        enriched.append(copy)
    return enriched


def _enrich_cc_cache_lines_skus(
    db: Session,
    warehouse_id: int,
    lines: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    if not lines:
        return lines

    missing_item_ids: list[int] = []
    for line in lines:
        sku = line.get("sku")
        if sku is not None and str(sku).strip() != "":
            continue
        raw_id = line.get("item_id")
        if raw_id is None:
            continue
        item_id = int(raw_id)
        if item_id not in missing_item_ids:
            missing_item_ids.append(item_id)

    sku_by_item_id: dict[int, str] = {}
    if missing_item_ids:
        rows = (
            db.query(Item.id, Item.sku)
            .filter(
                Item.id.in_(missing_item_ids),
                Item.warehouse_id == warehouse_id,
            )
            .all()
        )
        sku_by_item_id = {row.id: row.sku for row in rows if row.sku}

    enriched: list[dict[str, Any]] = []
    for line in lines:
        merged = dict(line)
        sku = merged.get("sku")
        if sku is None or str(sku).strip() == "":
            raw_id = merged.get("item_id")
            if raw_id is not None:
                resolved = sku_by_item_id.get(int(raw_id))
                if resolved:
                    merged["sku"] = resolved
        enriched.append(merged)
    return enriched


def _detail_to_cache_line(detail: OutboundOrderDetail) -> dict[str, Any]:
    meta = detail.details if isinstance(detail.details, dict) else {}
    return {
        "detail_id": detail.id,
        "outbound_order_id": detail.outbound_order_id,
        "item_id": detail.item_id,
        "quantity": detail.quantity,
        "unit": detail.unit,
        "detail_type": detail.detail_type,
        # BM.04 / Masan thường nằm trong JSON cột details
        "vehicle_no": meta.get("vehicle_no"),
        "customer_name": meta.get("customer_name"),
        "sku": _resolve_line_sku(detail, meta),
        "trip": meta.get("trip"),
        "nvt": meta.get("nvt"),
        "lot_number": format_lot_value_for_display(
            str(meta.get("lot_number") or "").strip() or None,
        ),
        "pallet_count": meta.get("pallet_count"),
        "locator": meta.get("locator"),
    }


def assign_cc_zone(db: Session, warehouse_id: int) -> None:                     
    keys = cache_scan_keys(f"outbound:assigned_cc_location:{warehouse_id}:*")
    existed_locations: list[str] = []
    for key in keys:
        suffix = key.rsplit(":", 1)[-1]
        if suffix.isdigit():
            existed_locations.append(int(suffix))

    logger.info(f"Existed locations: {existed_locations}")

    numeric_suffix = cast(
        func.nullif(
            func.regexp_replace(Location.location_name, r"\D", "", "g"), ""
        ),
        Integer,
    )
    available_locations = (
        db.query(Location)
        .join(Zone, Zone.id == Location.zone_id)
        .filter(
            Location.warehouse_id == warehouse_id,
            Location.is_active.is_(True),
            Zone.code.in_(settings.zone_cc),
            ~Location.id.in_(existed_locations),
        )
        .order_by(numeric_suffix.asc().nulls_last(), Location.location_name.asc())
        .all()
    )

    logger.info(f"Available locations: {available_locations}")

    if not available_locations:
        return
    
    _ensure_warehouse_exists(db, warehouse_id)
    available_orders = db.query(OutboundOrder).filter(
        OutboundOrder.status != "completed",
        OutboundOrder.warehouse_id == warehouse_id
    ).all()
    if not available_orders:
        return
    order_ids = [o.id for o in available_orders]

    keys = cache_scan_keys(f"outbound:assigned_cc_details:{warehouse_id}:*")
    assigned_ids: list[int] = []
    for key in keys:
        suffix = key.rsplit(":", 1)[-1]
        if suffix.isdigit():
            assigned_ids.append(int(suffix))

    list_details = (
        db.query(OutboundOrderDetail)
        .options(joinedload(OutboundOrderDetail.item))
        .filter(
            OutboundOrderDetail.outbound_order_id.in_(order_ids),
            OutboundOrderDetail.status == "initialize",
            ~OutboundOrderDetail.id.in_(assigned_ids),
        )
        .all()
    )

    details_by_vehicle: dict[str, list] = defaultdict(list)
    reversed_vehicles_list = []
    for detail in list_details:
        meta = detail.details if isinstance(detail.details, dict) else {}
        raw = meta.get("vehicle_no")
        vehicle_number = str(raw).strip() if raw is not None else ""
        if not vehicle_number:
            vehicle_number = "no_vehicle"
        line = _detail_to_cache_line(detail)
        details_by_vehicle[vehicle_number].append(line)
        if vehicle_number not in reversed_vehicles_list:
            reversed_vehicles_list.append(vehicle_number)

    i = 0
    while i < len(available_locations) and i < len(reversed_vehicles_list):
        key = f"outbound:assigned_cc_location:{warehouse_id}:{available_locations[i].id}"
        lines = details_by_vehicle[reversed_vehicles_list[i]]
        outbound_order_id = lines[0].get("outbound_order_id")
        logger.info(f"Outbound order ID: {outbound_order_id}")
        for line in lines:
            cache_data = {
                "zone": available_locations[i].zone.code,
                "location_id": available_locations[i].id,
            }
            cache_set(f"outbound:assigned_cc_details:{warehouse_id}:{line['detail_id']}", cache_data, -1)
        cache_set(key, {"outbound_order_id": outbound_order_id, "zone": available_locations[i].zone.code, "vehicle_number": reversed_vehicles_list[i], "lines": lines}, -1)
        i += 1

# def _assigning_algorithm(db: Session, list_details: list[OutboundOrderDetail]) -> None:
#     zone_ccs = settings.zone_cc

#     assign_cc_by_zone: dict[str, dict[str, list]] = defaultdict(lambda: defaultdict(list))
#     quantity_by_item_id: dict[int, int] = defaultdict(int)
#     details_by_item_id: dict[int, list] = defaultdict(list)

#     temp = {zone: 0 for zone in zone_ccs}

#     for detail in list_details:
#         quantity_by_item_id[detail.item_id] += detail.quantity

#         meta = detail.details if isinstance(detail.details, dict) else {}
#         raw = meta.get("vehicle_no")
#         vehicle_number = str(raw).strip() if raw is not None else ""
#         if not vehicle_number:
#             vehicle_number = "no_vehicle"
#         line = _detail_to_cache_line(detail)

#         details_by_item_id[detail.item_id][vehicle_number].append(line)

#     for item_id, quantity in quantity_by_item_id.items():
#         lowest_zone = min(temp, key=temp.get)
#         temp[lowest_zone] += quantity
#         for vehicle_number, lines in details_by_item_id[item_id].items():
#             assign_cc_by_zone[lowest_zone][vehicle_number].extend(lines)


def clear_completed_cc_assignments(
    db: Session, warehouse_id: int, *, order_id: int | None = None,
    persist_locations: bool = True,
) -> dict[str, int]:
    """Reconcile completed DB details with both CC indexes, including old buckets.

    The periodic caller has no pending DB writes. It commits recovered CC history
    before deleting cache. Callback callers use persist_locations=False; records
    still needing history recovery wait for the next periodic pass.
    """
    locations: dict[int, dict] = {}
    assignments: dict[int, dict] = {}
    candidates: dict[int, set[int]] = defaultdict(set)
    for scanned in cache_scan_keys(f"outbound:assigned_cc_location:{warehouse_id}:*"):
        suffix = scanned.rsplit(":", 1)[-1]
        if not suffix.isdigit():
            continue
        loc_id = int(suffix)
        bucket = cache_get(f"outbound:assigned_cc_location:{warehouse_id}:{loc_id}")
        if not isinstance(bucket, dict):
            continue
        locations[loc_id] = bucket
        for line in bucket.get("lines") or []:
            detail_id = str(line.get("detail_id", ""))
            if detail_id.isdigit():
                candidates[int(detail_id)].add(loc_id)
    for scanned in cache_scan_keys(f"outbound:assigned_cc_details:{warehouse_id}:*"):
        suffix = scanned.rsplit(":", 1)[-1]
        if not suffix.isdigit():
            continue
        detail_id = int(suffix)
        meta = cache_get(f"outbound:assigned_cc_details:{warehouse_id}:{detail_id}")
        if isinstance(meta, dict) and str(meta.get("location_id", "")).isdigit():
            assignments[detail_id] = meta
            candidates[detail_id].add(int(meta["location_id"]))
    result = {"locations_deleted": 0, "locations_updated": 0, "details_deleted": 0, "history_saved": 0}
    if not candidates:
        return result
    # Scalar columns bypass stale ORM status values after confirmation/callback.
    rows = db.query(
        OutboundOrderDetail.id, OutboundOrderDetail.outbound_order_id,
        OutboundOrderDetail.status, OutboundOrderDetail.sorting_location_id,
    ).join(OutboundOrder, OutboundOrder.id == OutboundOrderDetail.outbound_order_id).filter(
        OutboundOrder.warehouse_id == warehouse_id,
        OutboundOrderDetail.id.in_(candidates),
    ).all()
    order_by_detail = {row.id: row.outbound_order_id for row in rows}
    completed = {row.id for row in rows if row.status == "completed"
                 and (order_id is None or row.outbound_order_id == order_id)}
    valid_locations = {identifier for (identifier,) in db.query(Location.id).filter(
        Location.warehouse_id == warehouse_id,
        Location.id.in_({loc for ids in candidates.values() for loc in ids}),
    ).all()}
    for row in rows:
        if row.id not in completed or row.sorting_location_id is not None:
            continue
        known_locations = candidates[row.id] & valid_locations
        if not persist_locations or len(known_locations) != 1:
            completed.discard(row.id)
            continue
        result["history_saved"] += db.query(OutboundOrderDetail).filter(
            OutboundOrderDetail.id == row.id,
            OutboundOrderDetail.sorting_location_id.is_(None),
        ).update({OutboundOrderDetail.sorting_location_id: next(iter(known_locations))}, synchronize_session=False)
    if result["history_saved"]:
        db.commit()
    if not completed:
        return result
    for loc_id, bucket in locations.items():
        lines = bucket.get("lines") or []
        remaining = [line for line in lines if int(line.get("detail_id") or 0) not in completed]
        refs = {detail_id for detail_id, ids in candidates.items() if loc_id in ids}
        if not (refs & completed):
            continue
        replacement = None if not remaining and refs <= completed else {**bucket, "lines": remaining}
        if replacement is not None:
            replacement["outbound_order_id"] = next((order_by_detail[identifier]
                for identifier in sorted(refs - completed) if identifier in order_by_detail), bucket.get("outbound_order_id"))
        if cache_replace_if_unchanged(
            f"outbound:assigned_cc_location:{warehouse_id}:{loc_id}", bucket, replacement,
        ):
            result["locations_deleted" if replacement is None else "locations_updated"] += 1
    for detail_id, meta in assignments.items():
        if detail_id in completed and cache_replace_if_unchanged(
            f"outbound:assigned_cc_details:{warehouse_id}:{detail_id}", meta, None,
        ):
            result["details_deleted"] += 1
    if any(result.values()):
        logger.info("Cleared completed CC assignments warehouse_id=%s result=%s", warehouse_id, result)
    return result

def _clear_assigned_cc_location(warehouse_id: int, outbound_order_id: int) -> None:
    keys = cache_scan_keys(f"outbound:assigned_cc_location:{warehouse_id}:*")
    lines = []
    for key in keys:
        suffix = key.rsplit(":", 1)[-1]
        if suffix.isdigit():
            location_id = int(suffix)
            loc_key = f"outbound:assigned_cc_location:{warehouse_id}:{location_id}"
            bucket = cache_get(loc_key)
            if bucket and bucket.get("outbound_order_id") == outbound_order_id:
                lines.extend(bucket.get("lines") or [])
                cache_delete(loc_key)
                cache_delete_pattern(f"outbound:masan:stock_ready:{warehouse_id}:*:{location_id}")

    for line in lines:
        cache_delete(f"outbound:assigned_cc_details:{warehouse_id}:{line['detail_id']}")


def _clear_cc_cache_if_order_completed(
    db: Session,
    warehouse_id: int,
    order_id: int,
) -> None:
    order = (
        db.query(OutboundOrder)
        .filter(OutboundOrder.id == order_id)
        .first()
    )
    logger.info(f"Order: {order.id}, status: {order.status}")
    logger.info(f"Warehouse: {warehouse_id}")
    logger.info(f"Order ID: {order_id}")
    if order and order.status == "completed":
        _clear_assigned_cc_location(warehouse_id, order_id)

def get_sorting_data_for_zone(
    db: Session,
    warehouse_id: int,
    location_id: int,
) -> Optional[dict[str, Any]]:
    bucket = cache_get(f"outbound:assigned_cc_location:{warehouse_id}:{location_id}")
    if not bucket:
        return None

    raw_lines = bucket.get("lines") or []
    detail_ids = [
        int(line["detail_id"])
        for line in raw_lines
        if line.get("detail_id") is not None
    ]

    status_by_id: dict[int, str] = {}
    if detail_ids:
        detail_rows = (
            db.query(OutboundOrderDetail)
            .filter(OutboundOrderDetail.id.in_(detail_ids))
            .all()
        )
        status_by_id = {row.id: row.status for row in detail_rows}

    enriched_lines: list[dict[str, Any]] = []
    for line in raw_lines:
        merged = dict(line)
        detail_id = merged.get("detail_id")
        if detail_id is not None:
            merged["status"] = status_by_id.get(int(detail_id), "initialize")
        else:
            merged["status"] = None
        enriched_lines.append(merged)

    enriched_lines = _enrich_cc_cache_lines_skus(
        db, warehouse_id, enriched_lines
    )

    return {
        "zone": bucket.get("zone"),
        "vehicle_number": bucket.get("vehicle_number"),
        "lines": enriched_lines,
    }

# def get_item_needed_to_sorting(
#     db: Session, warehouse_id: int, zone: str
# ) -> dict[str, Any]:
#     keys = cache_scan_keys(f"outbound:assigned_cc_location:{warehouse_id}:*")
#     list_lines_raw: list[dict[str, Any]] = []
#     for key in keys:
#         suffix = key.rsplit(":", 1)[-1]
#         if suffix.isdigit():
#             location_id = int(suffix)
#             bucket = cache_get(f"outbound:assigned_cc_location:{warehouse_id}:{location_id}")
#             if bucket and bucket.get("zone") == zone:
#                 lines = bucket.get("lines") or []
#                 list_item_ids = [line.get("detail_id") for line in lines]
#                 list_lines_raw.extend(list_item_ids)

#     list_lines_raw = db.query(OutboundOrderDetail).filter(
#         OutboundOrderDetail.id.in_(list_lines_raw),
#     ).all()

#     detail_ids = [d.id for d in list_lines_raw]

#     list_lines = [line for line in list_lines_raw if line.unallocated_quantity > 0]
#     item_ids: list[int] = []
#     qty_by_item_id: dict[int, int] = defaultdict(int)
#     for line in list_lines:
#         item_id = line.item_id
#         qty_by_item_id[item_id] += int(line.unallocated_quantity)
#         if item_id not in item_ids:
#             item_ids.append(item_id)


#     if detail_ids:
#         allocations = (
#             db.query(OutboundOrderAllocation)
#             .filter(
#                 OutboundOrderAllocation.outbound_order_detail_id.in_(detail_ids),
#                 OutboundOrderAllocation.robot_task_id.isnot(None),
#             )
#             .order_by(OutboundOrderAllocation.id)
#             .all()
#         )
    
#         by_task: dict[int, list[OutboundOrderAllocation]] = defaultdict(list)
#         excluded_list = []
#         for alloc in allocations:
#             if alloc.outbound_order_detail_id in excluded_list:
#                 continue
#             if alloc.allocation_type == "outbound" and alloc.status not in ("completed", "failed", "cancelled"):
#                 excluded_list.append(alloc.outbound_order_detail_id)
#                 continue
#             if alloc.allocation_type == "return":
#                 by_task[alloc.outbound_order_detail_id].append(alloc)
#     else:
#         by_task = {}

#     if not item_ids:
#         return {
#             "warehouse_id": warehouse_id,
#             "zone": zone,
#             "items": [],
#             "line_count": len(list_lines),
#             "return_tasks": by_task,
#         }

#     item_rows = (
#         db.query(Item)
#         .filter(Item.id.in_(item_ids), Item.warehouse_id == warehouse_id)
#         .all()
#     )
#     item_by_id = {row.id: row for row in item_rows}

#     items_payload: list[dict[str, Any]] = []
#     for item_id in item_ids:
#         item = item_by_id.get(item_id)
#         if not item:
#             continue
#         items_payload.append(
#             {
#                 "item_id": item_id,
#                 "sku": item.sku,
#                 "item_name": item.name,
#                 "total_quantity": qty_by_item_id[item_id],
#             }
#         )

#     logger.info(f"Items payload: {items_payload}")
#     return {
#         "warehouse_id": warehouse_id,
#         "zone": zone,
#         "items": items_payload,
#         "line_count": len(list_lines),
#         "return_tasks": by_task,
#     }

def get_item_needed_to_sorting(
    db: Session, warehouse_id: int, zone: str
) -> dict[str, Any]:
    keys = cache_scan_keys(f"outbound:assigned_cc_location:{warehouse_id}:*")
    list_lines_raw: list[dict[str, Any]] = []
    for key in keys:
        suffix = key.rsplit(":", 1)[-1]
        if suffix.isdigit():
            location_id = int(suffix)
            bucket = cache_get(f"outbound:assigned_cc_location:{warehouse_id}:{location_id}")
            if bucket and bucket.get("zone") == zone:
                lines = bucket.get("lines") or []
                list_item_ids = [line.get("detail_id") for line in lines]
                list_lines_raw.extend(list_item_ids)

    list_lines_raw = db.query(OutboundOrderDetail).filter(
        OutboundOrderDetail.id.in_(list_lines_raw),
    ).all()

    detail_ids = [d.id for d in list_lines_raw]

    list_lines = [line for line in list_lines_raw if line.status != "completed"]
    item_ids: list[int] = []
    qty_by_item_id: dict[int, int] = defaultdict(int)
    for line in list_lines:
        item_id = line.item_id
        qty_by_item_id[item_id] += int(line.quantity or 0)
        if item_id not in item_ids:
            item_ids.append(item_id)


    if detail_ids:
        allocations = (
            db.query(OutboundOrderAllocation)
            .filter(
                OutboundOrderAllocation.outbound_order_detail_id.in_(detail_ids),
                OutboundOrderAllocation.robot_task_id.isnot(None),
            )
            .order_by(OutboundOrderAllocation.id)
            .all()
        )
    
        by_task: dict[int, list[OutboundOrderAllocation]] = defaultdict(list)
        excluded_list = []
        for alloc in allocations:
            if alloc.outbound_order_detail_id in excluded_list:
                continue
            if alloc.allocation_type == "outbound" and alloc.status != "completed":
                excluded_list.append(alloc.outbound_order_detail_id)
                continue
            if alloc.allocation_type == "return":
                by_task[alloc.outbound_order_detail_id].append(alloc)
    else:
        by_task = {}

    if not item_ids:
        return {
            "warehouse_id": warehouse_id,
            "zone": zone,
            "items": [],
            "line_count": len(list_lines),
            "return_tasks": by_task,
        }

    item_rows = (
        db.query(Item)
        .filter(Item.id.in_(item_ids), Item.warehouse_id == warehouse_id)
        .all()
    )
    item_by_id = {row.id: row for row in item_rows}

    items_payload: list[dict[str, Any]] = []
    for item_id in item_ids:
        item = item_by_id.get(item_id)
        if not item:
            continue
        items_payload.append(
            {
                "item_id": item_id,
                "sku": item.sku,
                "item_name": item.name,
                "total_quantity": qty_by_item_id[item_id],
            }
        )

    logger.info(f"Items payload: {items_payload}")
    return {
        "warehouse_id": warehouse_id,
        "zone": zone,
        "items": items_payload,
        "line_count": len(list_lines),
        "return_tasks": by_task,
    }

# def sending_masan_outbound_task(db: Session, warehouse_id: int, zone: str, item_id: int, to_location_id: int):
#     target = db.query(Location).filter(
#         Location.id == to_location_id,
#         Location.warehouse_id == warehouse_id,
#         Location.is_active.is_(True),
#     ).first()
#     if not target:
#         raise ValueError("Vị trí VT không tồn tại trong kho đang chọn")

#     keys = cache_scan_keys(f"outbound:assigned_cc_location:{warehouse_id}:*")
#     cached_detail_ids: set[int] = set()
#     for key in keys:
#         suffix = key.rsplit(":", 1)[-1]
#         if suffix.isdigit():
#             location_id = int(suffix)
#             bucket = cache_get(f"outbound:assigned_cc_location:{warehouse_id}:{location_id}")
#             if bucket and bucket.get("zone") == zone:
#                 lines = bucket.get("lines") or []
#                 for line in lines:
#                     if line.get("item_id") == item_id and line.get("detail_id") is not None:
#                         cached_detail_ids.add(int(line["detail_id"]))

#     # Keep retries for the same detail serialized, including partially picked lines.
#     details = (
#         db.query(OutboundOrderDetail)
#         .join(OutboundOrder, OutboundOrder.id == OutboundOrderDetail.outbound_order_id)
#         .filter(
#             OutboundOrderDetail.id.in_(cached_detail_ids),
#             OutboundOrderDetail.item_id == item_id,
#             OutboundOrder.warehouse_id == warehouse_id,
#         )
#         .order_by(OutboundOrderDetail.id)
#         .populate_existing()
#         .with_for_update(of=OutboundOrderDetail)
#         .all()
#     )
#     detail_ids = [detail.id for detail in details]
#     by_locator: dict[str, list[OutboundOrderDetail]] = defaultdict(list)
#     list_lines_db: list[OutboundOrderDetail] = []
#     for detail in details:
#         # Re-read after acquiring the lock: another call may just have allocated it.
#         db.refresh(detail, ["unallocated_quantity"])
#         if detail.unallocated_quantity <= 0:
#             continue
#         locator = str((detail.details or {}).get("locator") or "").strip()
#         if locator:
#             by_locator[locator].append(detail)
#         else:
#             list_lines_db.append(detail)

#     missing_locators: list[str] = []
#     for locator, locator_details in by_locator.items():
#         source = (
#             db.query(Location)
#             .filter(
#                 Location.warehouse_id == warehouse_id,
#                 Location.is_active.is_(True),
#                 or_(Location.location_name == locator, Location.location_code == locator),
#             )
#             .order_by(Location.id)
#             .with_for_update(of=Location)
#             .first()
#         )
#         stocks = []
#         if source:
#             stocks = (
#                 db.query(ItemStock)
#                 .filter(
#                     ItemStock.location_id == source.id,
#                     ItemStock.item_id == item_id,
#                     ItemStock.is_active.is_(True),
#                     ItemStock.status.in_(["available", "split"]),
#                     ItemStock.quantity > 0,
#                     countable_stock_level_criterion(),
#                 )
#                 .order_by(ItemStock.created_at, ItemStock.id)
#                 .populate_existing()
#                 .with_for_update(of=ItemStock)
#                 .all()
#             )
#             for stock in stocks:
#                 db.refresh(stock, ["available_quantity"])
#             stocks = [stock for stock in stocks if stock.available_quantity > 0]
#         if not stocks:
#             missing_locators.append(locator)
#             continue
#         lines = [
#             DetailForCalculate(
#                 id=detail.id,
#                 item_id=detail.item_id,
#                 quantity=int(detail.unallocated_quantity),
#                 detail_type=detail.detail_type,
#                 details=detail.details or {},
#             )
#             for detail in locator_details
#         ]
#         greedy_allocate_stocks_to_lines(
#             db, stocks, lines, locator_details[0].outbound_order_id, with_robot_task=True
#         )

#     if list_lines_db:
#         details_by_order: dict[int, list[OutboundOrderDetail]] = defaultdict(list)
#         for detail in list_lines_db:
#             details_by_order[detail.outbound_order_id].append(detail)
#         calculate_bodies: list[CalculateOutboundDetail] = []
#         for order_id, details in details_by_order.items():
#             line_items = [
#                 DetailForCalculate(
#                     id=d.id,
#                     item_id=d.item_id,
#                     quantity=int(d.unallocated_quantity),
#                     detail_type=d.detail_type,
#                     details=d.details if isinstance(d.details, dict) else {},
#                 )
#                 for d in details
#             ]
#             calculate_bodies.append(
#                 CalculateOutboundDetail(
#                     warehouse_id=warehouse_id,
#                     outbound_order_id=order_id,
#                     line_items=line_items,
#                 )
#             )
        
#         for body in calculate_bodies:
#             calculate_outbound_order(db, body, strategy="fefo")

#     db.flush()
#     picked_allocation = (
#         db.query(OutboundOrderAllocation)
#         .join(RobotTask, RobotTask.id == OutboundOrderAllocation.robot_task_id)
#         .options(joinedload(OutboundOrderAllocation.outbound_order_detail))
#         .filter(
#             OutboundOrderAllocation.outbound_order_detail_id.in_(detail_ids),
#             OutboundOrderAllocation.status == "initialize",
#             OutboundOrderAllocation.allocation_type == "outbound",
#             OutboundOrderAllocation.robot_task_id.isnot(None),
#             RobotTask.status == "initialize",
#         )
#         .order_by(OutboundOrderAllocation.id)
#         .first()
#     )

#     if not picked_allocation:
#         if missing_locators:
#             raise ValueError(
#                 f"Không có tồn khả dụng cho mã hàng tại Locator {', '.join(missing_locators)}. "
#                 "Hãy kiểm kê/cập nhật tồn đúng Locator trước khi gọi lại."
#             )
#         raise ValueError("Không có phần hàng mới để gửi; kiểm tra nhu cầu và các lệnh đang xử lý")

#     robot_task_id = picked_allocation.robot_task_id
#     _pack_and_execute_robot_task(db, robot_task_id, picked_allocation, to_location_id)
    
#     all_lacked = [
#         row.model_dump()
#         for order_id in sorted({detail.outbound_order_id for detail in details})
#         for row in lacked_details(db, order_id)
#         if row.id in cached_detail_ids
#     ]
#     return {
#         "warehouse_id": warehouse_id,
#         "zone": zone,
#         "item_id": item_id,
#         "lacked": all_lacked,
#     }    

def sending_masan_outbound_task(db: Session, warehouse_id: int, zone: str, item_id: int, to_location_id: int):
    from app.modules.masan.masan_auto_service import warehouse_lock, guard_manual_locations
    with warehouse_lock(db, warehouse_id):
        guard_manual_locations(db, warehouse_id, {to_location_id})
        return _sending_masan_outbound_task(db, warehouse_id, zone, item_id, to_location_id)


def _sending_masan_outbound_task(db: Session, warehouse_id: int, zone: str, item_id: int, to_location_id: int):
    from app.modules.masan.masan_auto_manual_service import context, prepare_dispatch
    if context(db) is not None:
        return prepare_dispatch(db, warehouse_id, zone, item_id, to_location_id)
    # Need optimize
    if _exist_robot_task(db, to_location_id):
        raise ValueError(f"Vị trí đích {to_location_id} đang có robot task chạy tới")

    keys = cache_scan_keys(f"outbound:assigned_cc_location:{warehouse_id}:*")
    list_lines: list[dict[str, Any]] = []
    has_lo_lines: list[dict[str, Any]] = []
    for key in keys:
        suffix = key.rsplit(":", 1)[-1]
        if suffix.isdigit():
            location_id = int(suffix)
            bucket = cache_get(f"outbound:assigned_cc_location:{warehouse_id}:{location_id}")
            if bucket and bucket.get("zone") == zone:
                lines = bucket.get("lines") or []
                for line in lines:
                    if line.get("item_id") == item_id:
                        if line.get("locator") is not None:
                            has_lo_lines.append(line["detail_id"])
                        else:
                            list_lines.append(line["detail_id"])

    list_lines_db = db.query(OutboundOrderDetail).filter(
        OutboundOrderDetail.id.in_(list_lines),
        OutboundOrderDetail.status == "initialize"
    ).all()

    has_lo_lines_db = db.query(OutboundOrderDetail).filter(
        OutboundOrderDetail.id.in_(has_lo_lines),
        OutboundOrderDetail.status == "initialize"
    ).all()

    results: list[dict[str, Any]] = []
    if len(has_lo_lines_db) > 0:
        for detail in has_lo_lines_db:
            locator = detail.details.get("locator")
            has_task = None
            if locator is not None:
                stock = (
                    db.query(ItemStock)
                    .join(ItemStock.location)
                    .filter(
                        Location.location_name == locator,
                        ItemStock.item_id == detail.item_id,
                        ItemStock.quantity >= detail.quantity,
                    )
                    .first()
                )
                allocation = (
                    db.query(OutboundOrderAllocation)
                    .join(OutboundOrderAllocation.from_location)
                    .filter(
                        OutboundOrderAllocation.status == "initialize",
                        OutboundOrderAllocation.allocation_type == "outbound",
                        Location.location_name == locator,
                    )
                    .first()
                )

                if allocation is not None:
                    has_task = allocation.robot_task_id
                
                if stock is not None:
                    if not has_task:
                        robot_task = RobotTask(
                            order_id=f"TDS_Outbound_{uuid.uuid4().hex[:8]}",
                            quantity=stock.quantity,
                            process_code=settings.outbound_process_code,
                            system_code="Thadosoft",
                            task_order_detail=json.dumps([{"taskPath": "None"}]),
                        )
                        db.add(robot_task)
                        db.flush()

                        has_task = robot_task.id

                    allocation = OutboundOrderAllocation(
                        outbound_order_detail_id=detail.id,
                        item_stock_id=stock.id,
                        quantity=detail.quantity,
                        status="initialize",
                        from_location_id=stock.location_id,
                        robot_task_id=has_task,
                        allocation_type="outbound",
                    )
                    db.add(allocation)
                    db.flush()

    if len(list_lines_db) > 0:
        details_by_order: dict[int, list[OutboundOrderDetail]] = defaultdict(list)
        for detail in list_lines_db:
            details_by_order[detail.outbound_order_id].append(detail)
        calculate_bodies: list[CalculateOutboundDetail] = []
        for order_id, details in details_by_order.items():
            line_items = [
                DetailForCalculate(
                    id=d.id,
                    item_id=d.item_id,
                    quantity=int(d.quantity),
                    detail_type=d.detail_type,
                    details=d.details if isinstance(d.details, dict) else {},
                )
                for d in details
            ]
            calculate_bodies.append(
                CalculateOutboundDetail(
                    warehouse_id=warehouse_id,
                    outbound_order_id=order_id,
                    line_items=line_items,
                )
            )
        
        for body in calculate_bodies:
            response_caculate = calculate_outbound_order(db, body, strategy="fefo")
            results.append({
                "outbound_order_id": response_caculate.outbound_order_id,
                "is_fully_allocated": response_caculate.is_fully_allocated,
                "lacked": [x.model_dump() for x in response_caculate.lacked],
            })

    # logger.info(f"Results: {list_lines}")

    picked_allocation = (
        db.query(OutboundOrderAllocation)
        .options(joinedload(OutboundOrderAllocation.outbound_order_detail))
        .filter(
            OutboundOrderAllocation.outbound_order_detail_id.in_(list_lines + has_lo_lines),
            OutboundOrderAllocation.status == "initialize",
            OutboundOrderAllocation.allocation_type == "outbound",
            OutboundOrderAllocation.robot_task_id.isnot(None),
        )
        .order_by(OutboundOrderAllocation.id)
        .first()
    )

    # logger.info(f"Picked allocation: {picked_allocation.model_dump()}")

    if not picked_allocation:
        return

    robot_task_id = picked_allocation.robot_task_id
    _pack_and_execute_robot_task(db, robot_task_id, picked_allocation, to_location_id)
    
    if results:
        all_lacked = [
            row
            for r in results
            for row in r["lacked"]
        ]
    else:
        all_lacked = []
    return {
        "warehouse_id": warehouse_id,
        "zone": zone,
        "item_id": item_id,
        "lacked": all_lacked,
    }    

def _pack_and_execute_robot_task(db: Session, robot_task_id: int, picked_allocation: OutboundOrderAllocation, to_location_id: int) -> None: 
    robot_task = db.query(RobotTask).filter(RobotTask.id == robot_task_id).first()
    if not robot_task:
        return

    task_allocations = (
        db.query(OutboundOrderAllocation)
        .filter(
            OutboundOrderAllocation.robot_task_id == robot_task_id,
            OutboundOrderAllocation.status == "initialize",
            OutboundOrderAllocation.allocation_type == "outbound",
        )
        .order_by(OutboundOrderAllocation.id)
        .all()
    )
    if not task_allocations:
        return
    
    detail_by_id = {
        a.outbound_order_detail_id: a.outbound_order_detail
        for a in task_allocations
        if a.outbound_order_detail is not None
    }

    for detail in detail_by_id.values():
        sum_quantity = sum(a.quantity for a in task_allocations if a.outbound_order_detail_id == detail.id)
        order = detail.outbound_order if detail is not None else None
        if detail is not None and order is not None:
            warehouse_id = order.warehouse_id
            cc_meta = cache_get(
                f"outbound:assigned_cc_details:{warehouse_id}:{detail.id}"
            )
            if isinstance(cc_meta, dict) and cc_meta.get("location_id") is not None:
                _apply_clear_outbound_masan_details(
                    db,
                    warehouse_id,
                    str(cc_meta.get("zone") or ""),
                    int(cc_meta["location_id"]),
                    detail.id,
                    sum_quantity,
                )
    
    from_location_id = picked_allocation.from_location_id
    if from_location_id is None:
        stock = picked_allocation.item_stock
        from_location_id = stock.location_id if stock else None
    if not from_location_id or not to_location_id:
        raise ValueError("From location or to location is not set")

    execute_body = OutboundRobotTaskCreate(
        order_id=robot_task.order_id,
        from_location_id=from_location_id,
        to_location_id=to_location_id,
        allocations=[
            AllocationOutboundTaskExecute(allocation_id=a.id)
            for a in task_allocations
        ],
    )
    execute_outbound_task(db, execute_body, detail_type="auto")

def _apply_clear_outbound_masan_details(
    db: Session,
    warehouse_id: int,
    zone: str,
    location_id: int,
    detail_id: int,
    quantity: int,
) -> None:
    detail_id = int(detail_id)
    location_id = int(location_id)
    zone = (zone or "").strip()

    loc_key = f"outbound:assigned_cc_location:{warehouse_id}:{location_id}"
    bucket = cache_get(loc_key)
    if not bucket:
        return

    lines = bucket.get("lines") or []
    new_lines = []
    changed = False
    for line in lines:
        if line.get("detail_id") is None or int(line["detail_id"]) != detail_id:
            if line.get("detail_id") is not None:
                new_lines.append(line)
            continue
        changed = True
        remaining = int(line["quantity"]) - quantity
        if remaining <= 0:
            remaining = 0
            # cache_delete(f"outbound:assigned_cc_details:{warehouse_id}:{detail_id}")
            continue
        new_lines.append({**line, "quantity": remaining})
    if not changed:
        return
    if new_lines:
        cache_set(loc_key, {**bucket, "lines": new_lines}, -1)
    # if not new_lines:
    #     cache_delete(loc_key)
    # else:
    #     cache_set(loc_key, {**bucket, "lines": new_lines}, -1)

def _clear_cache_details(db: Session, warehouse_id: int, detail: OutboundOrderDetail) -> None:
    detail_id = detail.id
    loc_key = f"outbound:assigned_cc_details:{warehouse_id}:{detail_id}"
    bucket = cache_get(loc_key)
    if not bucket:
        return

    location_id = bucket.get("location_id")
    if location_id is not None:
        detail.sorting_location_id = int(location_id)
        db.flush()
    cache_delete(f"outbound:assigned_cc_details:{warehouse_id}:{detail_id}")

    if location_id is None:
        return

    cc_key = f"outbound:assigned_cc_location:{warehouse_id}:{location_id}"
    bucket_cc = cache_get(cc_key)
    if not bucket_cc:
        return
    lines = bucket_cc.get("lines") or []
    new_lines = []
    for line in lines:
        if line.get("detail_id") is None or int(line["detail_id"]) != detail_id:
            if line.get("detail_id") is not None:
                new_lines.append(line)
            continue
    
    if new_lines:
        cache_set(cc_key, {**bucket_cc, "lines": new_lines}, -1)
    else:
        cache_delete(cc_key)


def clear_cc_location_caches(warehouse_id: int) -> dict[str, Any]:
    """
    Xóa toàn bộ Redis cache liên quan ô CC (assigned + stock_ready pending).
    Không xóa đơn/ detail trên DB.
    """
    if warehouse_id <= 0:
        raise ValueError("warehouse_id is required")

    assigned_location_keys = cache_delete_pattern(
        f"outbound:assigned_cc_location:{warehouse_id}:*"
    )
    assigned_details_keys = cache_delete_pattern(
        f"outbound:assigned_cc_details:{warehouse_id}:*"
    )
    stock_ready_keys = cache_delete_pattern(
        f"outbound:masan:stock_ready:{warehouse_id}:*"
    )
    total = assigned_location_keys + assigned_details_keys + stock_ready_keys
    return {
        "warehouse_id": warehouse_id,
        "zones": list(settings.zone_cc),
        "assigned_cc_location_keys": assigned_location_keys,
        "assigned_cc_details_keys": assigned_details_keys,
        "stock_ready_keys": stock_ready_keys,
        "total_keys_deleted": total,
    }


def get_sorting_zone_pending_stock(
    db: Session,
    warehouse_id: int,
    zone: str,
    *,
    republish: bool = False,
) -> dict[str, Any]:
    zone_norm = (zone or "").strip()
    if not zone_norm:
        raise ValueError("Zone is required")

    by_location: dict[int, list[dict[str, Any]]] = defaultdict(list)
    seen_keys: set[str] = set()
    pattern = f"outbound:masan:stock_ready:{warehouse_id}:{zone_norm}:*"

    for scanned in cache_scan_keys(pattern):
        parts = scanned.rsplit(":", 2)
        if len(parts) < 3:
            continue
        try:
            allocation_id = int(parts[-2])
            location_id = int(parts[-1])
        except ValueError:
            continue
        logical_key = (
            f"outbound:masan:stock_ready:{warehouse_id}:{zone_norm}:"
            f"{allocation_id}:{location_id}"
        )
        if logical_key in seen_keys:
            continue
        seen_keys.add(logical_key)
        payload = cache_get(logical_key)
        if not isinstance(payload, dict):
            continue
        by_location[location_id].append(payload)

    name_by_id: dict[int, str | None] = {}
    if by_location:
        for loc in (
            db.query(Location)
            .filter(
                Location.id.in_(list(by_location.keys())),
                Location.warehouse_id == warehouse_id,
            )
            .all()
        ):
            name_by_id[loc.id] = (loc.location_name or "").strip() or None

    locations_payload: list[dict[str, Any]] = []
    for loc_id in sorted(by_location.keys()):
        locations_payload.append(
            {
                "location_id": loc_id,
                "location_name": name_by_id.get(loc_id),
                "allocations": _enrich_stock_ready_payloads(db, by_location[loc_id]),
            }
        )

    published_events = 0
    if republish:
        detail_payloads: dict[int, list[dict[str, Any]]] = defaultdict(list)
        detail_location: dict[int, int] = {}
        for loc_id, rows in by_location.items():
            for row in rows:
                raw_detail = row.get("outbound_order_detail_id")
                if raw_detail is None:
                    continue
                detail_id = int(raw_detail)
                detail_payloads[detail_id].append(row)
                detail_location[detail_id] = loc_id
        for detail_id, payloads in detail_payloads.items():
            publish_masan_sorting_zone(
                warehouse_id,
                zone_norm,
                event_type=EVENT_MASAN_SORTING_STOCK_READY,
                data={
                    "location_id": detail_location[detail_id],
                    "detail_id": detail_id,
                    "details": payloads,
                },
            )
            published_events += 1

    return {
        "warehouse_id": warehouse_id,
        "zone": zone_norm,
        "locations": locations_payload,
        "published_events": published_events,
    }


def _delete_stock_ready_cache_for_allocation(
    warehouse_id: int,
    zone: str,
    allocation_id: int,
    *,
    status: str,
    quantity: int | None = None,
) -> None:
    # zone_norm = (zone or "").strip()
    # cache_delete_pattern(
    #     f"outbound:masan:stock_ready:{warehouse_id}:{zone_norm}:{allocation_id}:*"
    # )

    zone_norm = (zone or "").strip()
    pattern = f"outbound:masan:stock_ready:{warehouse_id}:{zone_norm}:{allocation_id}:*"
    updated = 0
    for scanned in cache_scan_keys(pattern):
        parts = scanned.rsplit(":", 2)
        if len(parts) < 3:
            continue
        try:
            loc_id = int(parts[-1])
        except ValueError:
            continue
        logical_key = (
            f"outbound:masan:stock_ready:{warehouse_id}:{zone_norm}:"
            f"{allocation_id}:{loc_id}"
        )
        payload = cache_get(logical_key)
        if not isinstance(payload, dict):
            continue
        payload = {**payload, "status": status}
        if quantity is not None:
            payload["quantity"] = int(quantity)
        cache_set(logical_key, payload, -1)
        updated += 1
    return updated


def confirm_allocation_outbound(
    db: Session,
    warehouse_id: int,
    zone: str,
    quantities: dict[int, int],
    sorting_persons: str,
) -> None:
    from app.modules.masan.masan_auto_service import confirm as confirm_automatic
    if confirm_automatic(db, warehouse_id, zone, quantities, sorting_persons):
        return

    from app.modules.masan.masan_auto_service import warehouse_lock, wake
    with warehouse_lock(db, warehouse_id):
        result = _confirm_allocation_outbound(db, warehouse_id, zone, quantities, sorting_persons)
    wake()
    return result


def _confirm_allocation_outbound(db: Session, warehouse_id: int, zone: str,
                                 quantities: dict[int, int], sorting_persons: str) -> None:

    allocation_ids = list(quantities.keys())
    if not allocation_ids:
        raise ValueError("quantities must contain at least one allocation")

    target_allocations = (
        db.query(OutboundOrderAllocation)
        .filter(OutboundOrderAllocation.id.in_(allocation_ids))
        .all()
    )

    if len(target_allocations) != len(allocation_ids):
        found = {a.id for a in target_allocations}
        missing = sorted(set(allocation_ids) - found)
        raise ValueError(f"Allocations not found: {missing}")

    stocks: list[ItemStock] = []
    outbound_order = None
    for allocation in target_allocations:
        allocation.status = "confirmed"
        allocation.sorting_persons = sorting_persons
        _delete_stock_ready_cache_for_allocation(
                warehouse_id, zone, allocation.id, status="completed", quantity=quantities[allocation.id]
            )
        if allocation.end_sorting_time is None:
            allocation.end_sorting_time = datetime.now()

        db.flush()

        detail = allocation.outbound_order_detail
        still_pending = (
            db.query(OutboundOrderAllocation.id)
            .filter(
                OutboundOrderAllocation.outbound_order_detail_id == detail.id,
                OutboundOrderAllocation.allocation_type == "outbound",
                OutboundOrderAllocation.status.notin_(["completed", "confirmed", "failed", "cancelled"]),
            )
            .first()
        )

        if quantities[allocation.id] != allocation.quantity:
            allocation.quantity = quantities[allocation.id]
            allocation.outbound_order_detail.quantity = allocation.outbound_order_detail.quantity - quantities[allocation.id]
            if allocation.item_stock not in stocks:
                stocks.append(allocation.item_stock)

        db.commit()
        db.refresh(allocation.outbound_order_detail)

        # db.refresh(detail, ["unallocated_quantity"])
        # if not still_pending and detail.unallocated_quantity <= 0:
        if not still_pending:
            _clear_cache_details(db, warehouse_id, detail)
            outbound_order = detail.outbound_order
            
    _settle_outbound_stock(db, target_allocations)

    if outbound_order:
        _clear_cc_cache_if_order_completed(db, warehouse_id, outbound_order.id)


    for stock in stocks:
        return_allocation = db.query(OutboundOrderAllocation).filter(
            OutboundOrderAllocation.item_stock_id == stock.id,
            OutboundOrderAllocation.allocation_type == "return",
        ).first()

        if not return_allocation:
            reverse_allocation = db.query(OutboundOrderAllocation).filter(
                OutboundOrderAllocation.item_stock_id == stock.id,
                OutboundOrderAllocation.allocation_type == "outbound",
            ).first()
            order_id = f"TDS_Return_{uuid.uuid4().hex[:8]}"
            return_task = RobotTask(
                order_id=order_id,
                quantity=stock.quantity,
                process_code=settings.inbound_process_code,
                system_code="Thadosoft",
                task_order_detail=json.dumps([{"taskPath": f"{reverse_allocation.to_location.location_code},{reverse_allocation.from_location.location_code}"}])
            )
            db.add(return_task)
            db.flush()
            return_allocation = OutboundOrderAllocation(
                outbound_order_detail_id=reverse_allocation.outbound_order_detail_id,
                item_stock_id=stock.id,
                quantity=stock.quantity,
                status="initialize",
                from_location_id=reverse_allocation.to_location_id,
                to_location_id=reverse_allocation.from_location_id,
                robot_task_id=return_task.id,
                allocation_type="return",
            )
            db.add(return_allocation)
            from app.modules.masan.masan_return_journal import register
            register(db, return_task)
            

        else:
            return_allocation.quantity = stock.quantity
            
        db.flush()

    # Persist newly created return allocations before reconciling completion.
    # db.commit()
    # clear_completed_cc_assignments(db, warehouse_id)

def _send_taking_stock(db: Session, robot_task_id: int, allocations: list[OutboundOrderAllocation]) -> None:
    logger.info(f"Sending taking stock for robot task {robot_task_id}")
    
    sending_payloads: dict[int, list[dict[str, Any]]] = defaultdict(list)
    for allocation in allocations:
        if allocation.allocation_type != "outbound":
            return
        detail = allocation.outbound_order_detail
        if not detail:
            continue
        
        warehouse_id = detail.outbound_order.warehouse_id
        sending_payloads[detail.id].append({
            "id": allocation.id,
            "outbound_order_detail_id": allocation.outbound_order_detail_id,
            "item_stock_id": allocation.item_stock_id,
            "quantity": int(allocation.quantity),
            "status": allocation.status,
            "allocation_type": allocation.allocation_type,
            "robot_task_id": allocation.robot_task_id,
            "from_location_id": allocation.from_location_id,
            "to_location_id": allocation.to_location_id,
        })

    for detail_id, payloads in sending_payloads.items():
        payloads = _enrich_stock_ready_payloads(db, payloads)
        sending_payloads[detail_id] = payloads
        cache_data = cache_get(f"outbound:assigned_cc_details:{warehouse_id}:{detail_id}")
        if cache_data:
            location_id = cache_data.get("location_id")
            zone = cache_data.get("zone")

            publish_masan_sorting_zone(
                warehouse_id,
                zone,
                event_type=EVENT_MASAN_SORTING_STOCK_READY,
                data={
                    "location_id": location_id,
                    "details": payloads,
                },
            )

            logger.info(f"Publishing sorting stock ready for warehouse_id={warehouse_id}, zone={zone}, location_id={location_id}, details={len(payloads)}")

            for payload in payloads:
                key = (
                    f"outbound:masan:stock_ready:{warehouse_id}:{zone}:"
                    f"{payload['id']}:{location_id}"
                )
                cache_set(key, payload, -1)


def list_sorting_zone_cc_locations(
    db: Session,
    warehouse_id: int,
    zone: str,
) -> dict[str, Any]:
    """Trạng thái gán CC theo mã zone (Zone.code) — mọi location trong zone DB."""
    zone_norm = (zone or "").strip()
    if not zone_norm:
        raise ValueError("Zone is required")
    _ensure_warehouse_exists(db, warehouse_id)

    zone_row = (
        db.query(Zone)
        .filter(Zone.warehouse_id == warehouse_id, Zone.code == zone_norm)
        .first()
    )
    if not zone_row:
        return {
            "warehouse_id": warehouse_id,
            "zone": zone_norm,
            "locations": [],
        }

    location_rows = (
        db.query(Location)
        .filter(
            Location.warehouse_id == warehouse_id,
            Location.zone_id == zone_row.id,
        )
        .order_by(Location.location_code)
        .all()
    )

    items: list[dict[str, Any]] = []
    for loc in location_rows:
        location_name = (loc.location_name or "").strip() or None
        bucket_data = get_sorting_data_for_zone(db, warehouse_id, loc.id)
        if not bucket_data:
            items.append(
                {
                    "location_id": loc.id,
                    "location_code": loc.location_code,
                    "location_name": location_name,
                    "assigned": False,
                    "zone": zone_norm,
                    "vehicle_number": None,
                    "lines": [],
                }
            )
            continue
        items.append(
            {
                "location_id": loc.id,
                "location_code": loc.location_code,
                "location_name": location_name,
                "assigned": True,
                "zone": bucket_data.get("zone"),
                "vehicle_number": bucket_data.get("vehicle_number"),
                "lines": bucket_data.get("lines") or [],
            }
        )

    return {
        "warehouse_id": warehouse_id,
        "zone": zone_norm,
        "locations": items,
    }

def start_sorting_time(db: Session, allocation_ids: list[int]) -> dict[str, Any]:
    if not allocation_ids:
        return {"started_at": None, "already_started": False}

    allocations = (
        db.query(OutboundOrderAllocation)
        .filter(
            OutboundOrderAllocation.id.in_(allocation_ids),
            OutboundOrderAllocation.allocation_type == "outbound",
        )
        .all()
    )
    if not allocations:
        return {"started_at": None, "already_started": False}

    existing = [
        row.start_sorting_time
        for row in allocations
        if row.start_sorting_time is not None
    ]
    if existing:
        return {"started_at": min(existing), "already_started": True}

    started_at = datetime.now()
    for allocation in allocations:
        allocation.start_sorting_time = started_at
    db.commit()
    return {"started_at": started_at, "already_started": False}



#Fix bấm vị trí gọi lệnh vẫn ra chỗ đã được chọn - NG
#Fix Chọn được vị trí cho đơn xuất.
#Fix Bypass
#Fix export đơn nhập ra theo thời gian
#Fix gọi khi click confirm 1 click cho tất cả tại vị trí VT - OK
#Fix update detail để hoàn thành được lệnh nếu user gửi số lượng lấy ít hơn số có - OK
#Add lệnh return nếu sửa đổi số lượng lấy ra trong đơn - OK
#Fix gọi robot tự động.
#Fix không cho gửi nếu vị trí đó có hàng.
#Fix QR code đơn nhập
#Fix import phiên ra qrcode.

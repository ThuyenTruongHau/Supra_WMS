"""Shared Excel layout for Masan outbound import (BM.04) and export report."""

from __future__ import annotations

from typing import TYPE_CHECKING, Any

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

EXPORT_TITLE = (
    "BIỂU MẪU LIST XE XUẤT HÀNG\n( Cung cấp danh sách lấy hàng theo xe )"
)

EXPORT_HEADERS: tuple[str, ...] = (
    "Số xe",
    "Tên khách hàng",
    "Trip",
    "NVT",
    "Mã Item",
    "Tên Item",
    "LOT",
    "Lot status",
    "Số lượng",
    "Số pallet",
    "Locator",
)

OUTBOUND_REPORT_TITLE = (
    "BIỂU MẪU BÁO CÁO XUẤT HÀNG\n(New System trả về để truy vết)"
)

OUTBOUND_REPORT_EXTRA_HEADERS: tuple[str, ...] = (
    "Người chia",
    "Quãng đường (m)",
    "Robot",
    "Thời gian bắt đầu",
    "Thời gian kết thúc",
    "Pin bắt đầu",
    "Pin kết thúc",
    "Vị trí chia",
    "Bắt đầu chia",
    "Kết thúc chia",
)

OUTBOUND_REPORT_HEADERS: tuple[str, ...] = (
    *EXPORT_HEADERS,
    "Vị trí xuất",
    *OUTBOUND_REPORT_EXTRA_HEADERS,
)

HEADER_ROW = 2
DATA_START_ROW = 3

OUTBOUND_REPORT_COLUMN_WIDTHS: tuple[float, ...] = (
    12.0,  # Số xe
    28.0,  # Tên khách hàng
    10.0,  # Trip
    12.0,  # NVT
    12.0,  # Mã Item
    40.0,  # Tên Item
    10.0,  # LOT
    12.0,  # Lot status
    10.0,  # Số lượng
    10.0,  # Số pallet
    14.0,  # Locator
    16.0,  # Vị trí xuất
    18.0,  # Người chia
    12.0,  # Quãng đường (m)
    10.0,  # Robot
    18.0,  # Thời gian bắt đầu
    18.0,  # Thời gian kết thúc
    12.0,  # Pin bắt đầu
    12.0,  # Pin kết thúc
    16.0,  # Vị trí chia
    18.0,  # Bắt đầu chia
    18.0,  # Kết thúc chia
)

TITLE_ROW_HEIGHT = 42.0
HEADER_ROW_HEIGHT = 28.0
DATA_ROW_HEIGHT = 20.0

# 1-based column indexes for outbound report sheet.
OUTBOUND_REPORT_LOT_COLUMN = 7
OUTBOUND_REPORT_QUANTITY_COLUMN = 9
OUTBOUND_REPORT_DISTANCE_COLUMN = 14
OUTBOUND_REPORT_TEXT_FORMAT_COLUMNS = frozenset({OUTBOUND_REPORT_LOT_COLUMN})
OUTBOUND_REPORT_WRAP_TEXT_COLUMNS = frozenset({6, 12, 13, 16, 17, 20, 21, 22})

TRIP_PLACEHOLDERS = frozenset({"…", "...", ".", "-", "—"})

COLUMN_ALIASES: dict[str, list[str]] = {
    "vehicle_no": ["số xe"],
    "customer_name": ["tên khách hàng"],
    "trip": ["trip"],
    "nvt": ["nvt"],
    "sku": ["mã item"],
    "item_name": ["tên item"],
    "lot": ["lot"],
    "lot_status": ["lot status"],
    "quantity": ["số lượng"],
    "pallet_count": ["số pallet"],
    "locator": ["locator"],
}

_TITLE_FILL = PatternFill("solid", fgColor="FFFF00")
_HEADER_FONT = Font(bold=True, size=10)
_TITLE_FONT = Font(bold=True, size=11)
_DATA_FONT = Font(size=10)
_THIN_BORDER = Border(
    left=Side(style="thin", color="000000"),
    right=Side(style="thin", color="000000"),
    top=Side(style="thin", color="000000"),
    bottom=Side(style="thin", color="000000"),
)
_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
_LEFT_TOP_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)


def normalize_outbound_report_cell_value(column_index: int, value: Any) -> Any:
    if value is None or value == "":
        return None
    if column_index in OUTBOUND_REPORT_TEXT_FORMAT_COLUMNS:
        return str(value)
    return value


def apply_outbound_report_cell_style(ws: Worksheet, row: int, column: int) -> None:
    cell = ws.cell(row=row, column=column)
    cell.font = _DATA_FONT
    cell.border = _THIN_BORDER
    if column in OUTBOUND_REPORT_TEXT_FORMAT_COLUMNS:
        cell.number_format = "@"
    if column in OUTBOUND_REPORT_WRAP_TEXT_COLUMNS:
        cell.alignment = _LEFT_TOP_WRAP
    elif column in (OUTBOUND_REPORT_QUANTITY_COLUMN, OUTBOUND_REPORT_DISTANCE_COLUMN):
        cell.alignment = Alignment(horizontal="right", vertical="center")
    else:
        cell.alignment = _LEFT_WRAP


def style_outbound_report_worksheet(ws: Worksheet, *, last_data_row: int) -> None:
    """BM.04 columns + analysis fields; yellow title row, grid borders."""
    col_count = len(OUTBOUND_REPORT_HEADERS)

    for col_idx, width in enumerate(OUTBOUND_REPORT_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for col_idx in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _TITLE_FILL

    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=col_count)
    title_cell = ws.cell(row=1, column=2)
    title_cell.value = OUTBOUND_REPORT_TITLE
    title_cell.font = _TITLE_FONT
    title_cell.alignment = _CENTER_WRAP
    ws.row_dimensions[1].height = TITLE_ROW_HEIGHT

    for col_idx in range(1, col_count + 1):
        header_cell = ws.cell(row=HEADER_ROW, column=col_idx)
        header_cell.font = _HEADER_FONT
        header_cell.alignment = _CENTER_WRAP
        header_cell.border = _THIN_BORDER
    ws.row_dimensions[HEADER_ROW].height = HEADER_ROW_HEIGHT

    for row_idx in range(DATA_START_ROW, last_data_row + 1):
        ws.row_dimensions[row_idx].height = DATA_ROW_HEIGHT
        for col_idx in range(1, col_count + 1):
            apply_outbound_report_cell_style(ws, row_idx, col_idx)

    ws.freeze_panes = ws.cell(row=DATA_START_ROW, column=1)
